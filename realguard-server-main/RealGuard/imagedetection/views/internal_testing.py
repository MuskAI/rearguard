"""Persistent, administrator-only model evaluation and load-test support."""
from __future__ import annotations

import hashlib
import io
import ipaddress
import json
import math
import mimetypes
import os
import re
import resource
import shutil
import socket
import sqlite3
import statistics
import threading
import time
import uuid
import zipfile
from concurrent.futures import FIRST_COMPLETED, ThreadPoolExecutor, as_completed, wait
from collections import Counter
from datetime import datetime, timedelta, timezone
from pathlib import Path, PurePosixPath
from urllib.parse import urljoin, urlparse

import requests
from PIL import Image, UnidentifiedImageError

from imagedetection import legal_documents


DATA_ROOT = Path(
    os.environ.get("REALGUARD_INTERNAL_TEST_ROOT", "/opt/realguard-data/internal-testing")
)
IMPORT_ROOT = DATA_ROOT / "imports"
DB_PATH = Path(
    os.environ.get("REALGUARD_INTERNAL_TEST_DB", str(DATA_ROOT / "internal-testing.sqlite3"))
)
MAX_UPLOAD_BYTES = 24 * 1024 * 1024
MAX_IMPORT_CHUNK_BYTES = 8 * 1024 * 1024
# A value of zero means there is no count-based ingestion limit. Byte and storage
# boundaries remain in place so large datasets fail predictably instead of exhausting RAM.
MAX_DATASET_SAMPLES = max(0, int(os.environ.get("REALGUARD_INTERNAL_TEST_MAX_SAMPLES", "0")))
MAX_WEB_IMAGES = max(0, int(os.environ.get("REALGUARD_INTERNAL_TEST_MAX_WEB_IMAGES", "0")))
MAX_UPLOAD_FILES = max(0, int(os.environ.get("REALGUARD_INTERNAL_TEST_MAX_UPLOAD_FILES", "0")))
MAX_REDIRECTS = 4
MAX_EVALUATION_CONCURRENCY = 4
MAX_LOAD_CONCURRENCY = 16
MAX_LOAD_REQUESTS = 1000
MAX_LOAD_DURATION_SECONDS = 120
MAX_STORED_DATASETS = max(0, int(os.environ.get("REALGUARD_INTERNAL_TEST_MAX_DATASETS", "0")))
MIN_FREE_STORAGE_BYTES = max(
    0,
    int(os.environ.get("REALGUARD_INTERNAL_TEST_MIN_FREE_BYTES", str(2 * 1024 * 1024 * 1024))),
)
ALLOWED_IMAGE_FORMATS = {"JPEG", "PNG", "WEBP", "BMP", "GIF", "TIFF", "HEIF", "HEIC"}
ALLOWED_LABELS = {"real", "fake", "unlabeled"}
REAL_PATH_LABELS = {
    "real", "reals", "authentic", "genuine", "natural", "camera", "captured",
    "photograph", "photo", "original", "human", "0",
    "真实", "真图", "实拍", "自然图像", "相机", "原图",
}
FAKE_PATH_LABELS = {
    "fake", "fakes", "synthetic", "generated", "generation", "ai", "aigc",
    "ai_generated", "deepfake", "diffusion", "gan", "sdxl", "stable_diffusion", "midjourney",
    "dalle", "flux", "1", "生成", "生成图", "假图", "合成图", "人工合成",
    "即梦", "豆包",
}
_SCHEMA_LOCK = threading.Lock()
_IMPORT_LOCK = threading.RLock()
_SCHEMA_READY = False
_EXECUTOR = ThreadPoolExecutor(max_workers=4, thread_name_prefix="internal-testing")
_IMPORT_DETECTION_EXECUTOR = ThreadPoolExecutor(
    max_workers=4,
    thread_name_prefix="internal-testing-upload-detection",
)
_ACTIVE_IMPORTS: set[str] = set()
_ACTIVE_IMPORT_DETECTIONS: set[str] = set()


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _connect() -> sqlite3.Connection:
    ensure_schema()
    connection = sqlite3.connect(DB_PATH, timeout=30)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA foreign_keys = ON")
    connection.execute("PRAGMA busy_timeout = 30000")
    return connection


def ensure_schema() -> None:
    global _SCHEMA_READY
    if _SCHEMA_READY:
        return
    with _SCHEMA_LOCK:
        if _SCHEMA_READY:
            return
        DATA_ROOT.mkdir(parents=True, exist_ok=True)
        os.chmod(DATA_ROOT, 0o700)
        connection = sqlite3.connect(DB_PATH, timeout=30)
        try:
            connection.executescript(
                """
                PRAGMA journal_mode = WAL;
                PRAGMA foreign_keys = ON;
                CREATE TABLE IF NOT EXISTS datasets (
                    id TEXT PRIMARY KEY,
                    name TEXT NOT NULL,
                    source_type TEXT NOT NULL,
                    source_name TEXT,
                    default_label TEXT NOT NULL,
                    sample_count INTEGER NOT NULL DEFAULT 0,
                    labeled_count INTEGER NOT NULL DEFAULT 0,
                    total_bytes INTEGER NOT NULL DEFAULT 0,
                    created_at TEXT NOT NULL,
                    actor_id TEXT,
                    actor_name TEXT,
                    profile_name TEXT NOT NULL DEFAULT 'generic',
                    taxonomy_json TEXT NOT NULL DEFAULT '{}'
                );
                CREATE TABLE IF NOT EXISTS samples (
                    id TEXT PRIMARY KEY,
                    dataset_id TEXT NOT NULL,
                    name TEXT NOT NULL,
                    source TEXT,
                    sha256 TEXT NOT NULL,
                    mime_type TEXT NOT NULL,
                    width INTEGER NOT NULL,
                    height INTEGER NOT NULL,
                    byte_size INTEGER NOT NULL,
                    ground_truth TEXT NOT NULL,
                    storage_path TEXT NOT NULL,
                    relative_path TEXT,
                    label_source TEXT NOT NULL DEFAULT 'default',
                    class_path TEXT,
                    subclasses_json TEXT NOT NULL DEFAULT '{}',
                    group_id TEXT,
                    created_at TEXT NOT NULL,
                    FOREIGN KEY(dataset_id) REFERENCES datasets(id) ON DELETE CASCADE,
                    UNIQUE(dataset_id, sha256)
                );
                CREATE INDEX IF NOT EXISTS idx_samples_dataset ON samples(dataset_id);
                CREATE TABLE IF NOT EXISTS runs (
                    id TEXT PRIMARY KEY,
                    kind TEXT NOT NULL,
                    dataset_id TEXT,
                    model_id TEXT NOT NULL,
                    model_name TEXT,
                    status TEXT NOT NULL,
                    configuration_json TEXT NOT NULL,
                    metrics_json TEXT,
                    error TEXT,
                    completed_count INTEGER NOT NULL DEFAULT 0,
                    total_count INTEGER NOT NULL DEFAULT 0,
                    created_at TEXT NOT NULL,
                    started_at TEXT,
                    finished_at TEXT,
                    updated_at TEXT NOT NULL,
                    actor_id TEXT,
                    actor_name TEXT,
                    FOREIGN KEY(dataset_id) REFERENCES datasets(id) ON DELETE SET NULL
                );
                CREATE INDEX IF NOT EXISTS idx_runs_created ON runs(created_at DESC);
                CREATE TABLE IF NOT EXISTS results (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    run_id TEXT NOT NULL,
                    sample_id TEXT,
                    status TEXT NOT NULL,
                    predicted_label TEXT,
                    score REAL,
                    latency_ms INTEGER,
                    http_status INTEGER,
                    error TEXT,
                    response_json TEXT,
                    created_at TEXT NOT NULL,
                    FOREIGN KEY(run_id) REFERENCES runs(id) ON DELETE CASCADE,
                    FOREIGN KEY(sample_id) REFERENCES samples(id) ON DELETE SET NULL
                );
                CREATE INDEX IF NOT EXISTS idx_results_run ON results(run_id, id);
                """
            )
            sample_columns = {
                str(row[1]) for row in connection.execute("PRAGMA table_info(samples)").fetchall()
            }
            if "relative_path" not in sample_columns:
                connection.execute("ALTER TABLE samples ADD COLUMN relative_path TEXT")
            if "label_source" not in sample_columns:
                connection.execute(
                    "ALTER TABLE samples ADD COLUMN label_source TEXT NOT NULL DEFAULT 'default'"
                )
            if "class_path" not in sample_columns:
                connection.execute("ALTER TABLE samples ADD COLUMN class_path TEXT")
            if "subclasses_json" not in sample_columns:
                connection.execute(
                    "ALTER TABLE samples ADD COLUMN subclasses_json TEXT NOT NULL DEFAULT '{}'"
                )
            if "group_id" not in sample_columns:
                connection.execute("ALTER TABLE samples ADD COLUMN group_id TEXT")
            dataset_columns = {
                str(row[1]) for row in connection.execute("PRAGMA table_info(datasets)").fetchall()
            }
            if "profile_name" not in dataset_columns:
                connection.execute(
                    "ALTER TABLE datasets ADD COLUMN profile_name TEXT NOT NULL DEFAULT 'generic'"
                )
            if "taxonomy_json" not in dataset_columns:
                connection.execute(
                    "ALTER TABLE datasets ADD COLUMN taxonomy_json TEXT NOT NULL DEFAULT '{}'"
                )
            connection.commit()
            os.chmod(DB_PATH, 0o600)
        finally:
            connection.close()
        _SCHEMA_READY = True


def _row_dict(row: sqlite3.Row | None) -> dict | None:
    if row is None:
        return None
    payload = dict(row)
    for key in (
        "configuration_json", "metrics_json", "response_json", "taxonomy_json",
        "subclasses_json",
    ):
        if key not in payload:
            continue
        target = key.removesuffix("_json")
        try:
            payload[target] = json.loads(payload.pop(key) or "{}")
        except json.JSONDecodeError:
            payload[target] = {}
            payload.pop(key, None)
    return payload


def _actor_fields(actor: dict | None) -> tuple[str, str]:
    actor = actor or {}
    return (
        str(actor.get("adminId") or actor.get("Userid") or "")[:80],
        str(actor.get("username") or "")[:80],
    )


def _safe_name(value: str, fallback: str = "sample") -> str:
    name = Path(str(value or "")).name.strip().replace("\x00", "")
    return (name or fallback)[:180]


def _safe_relative_path(value: str, fallback: str = "sample") -> str:
    raw = str(value or "").replace("\\", "/").replace("\x00", "").strip()
    parts = [part.strip() for part in PurePosixPath(raw).parts if part not in ("", ".", "/")]
    if not parts or any(part == ".." for part in parts):
        return _safe_name(fallback)
    cleaned = [re.sub(r"[\x00-\x1f]", "", part)[:180] for part in parts]
    return "/".join(part for part in cleaned if part)[:1000] or _safe_name(fallback)


def _limit_reached(count: int, limit: int) -> bool:
    return bool(limit > 0 and count >= limit)


def _source_upload_limit(filename: str) -> int:
    return 0 if Path(str(filename or "")).suffix.lower() == ".zip" else MAX_UPLOAD_BYTES


def available_storage_bytes() -> int:
    DATA_ROOT.mkdir(parents=True, exist_ok=True)
    return max(0, int(shutil.disk_usage(DATA_ROOT).free) - MIN_FREE_STORAGE_BYTES)


def _ensure_storage_capacity(required_bytes: int) -> None:
    if max(0, int(required_bytes or 0)) > available_storage_bytes():
        raise ValueError("服务器磁盘剩余空间不足，无法继续导入该数据集")


def _source_stream(source):
    if isinstance(source, Path):
        return source.open("rb")
    stream = io.BytesIO(source) if isinstance(source, (bytes, bytearray)) else source
    try:
        stream.seek(0)
    except (AttributeError, OSError):
        pass
    return stream


def _read_source(source, limit: int = MAX_UPLOAD_BYTES) -> bytes:
    if isinstance(source, Path):
        with source.open("rb") as stream:
            data = stream.read(limit + 1)
    else:
        stream = _source_stream(source)
        data = stream.read(limit + 1)
    if len(data) > limit:
        raise ValueError("单个图片或文档超过 24 MB")
    return data


def _path_label(relative_path: str) -> tuple[str, str]:
    path = _safe_relative_path(relative_path)
    parts = list(PurePosixPath(path).parts[:-1])
    for part in reversed(parts):
        normalized = str(part).strip().lower().replace("-", "_").replace(" ", "_")
        tokens = {token for token in re.split(r"[^0-9a-z_\u4e00-\u9fff]+|_+", normalized) if token}
        candidates = tokens | {normalized}
        real = bool(candidates & REAL_PATH_LABELS)
        fake = bool(candidates & FAKE_PATH_LABELS)
        if real != fake:
            label = "real" if real else "fake"
            return label, f"directory:{part}"
    return "unlabeled", "unresolved"


def _path_parts(relative_path: str) -> list[str]:
    return list(PurePosixPath(_safe_relative_path(relative_path)).parts[:-1])


def _detect_dataset_profile(relative_paths: list[str]) -> str:
    """Recognize known benchmarks without assuming their selected root depth."""
    paths = [[part.lower() for part in _path_parts(path)] for path in relative_paths]
    if any("deepfake" in parts for parts in paths) and any(
        {"positive", "negative"} & set(parts) for parts in paths
    ):
        return "fraudbench"
    transforms = {"original", "redigital", "transfer"}
    for parts in paths:
        for index, part in enumerate(parts[:-1]):
            if part in transforms and parts[index + 1] in {"real", "ai"}:
                return "rrdataset"
    return "generic"


def _review_part(parts: list[str]) -> str:
    return next((part for part in parts if re.fullmatch(r"Review_\d+", part, re.I)), "")


def _classify_path(relative_path: str, profile: str) -> dict:
    """Return binary truth plus lossless, semantic directory dimensions."""
    safe_path = _safe_relative_path(relative_path)
    parts = _path_parts(safe_path)
    class_path = "/".join(parts)
    subclasses: dict[str, str] = {
        f"level_{index}": part for index, part in enumerate(parts, 1)
    }
    ground_truth = "unlabeled"
    label_source = "unresolved"
    group_id = ""

    if profile == "fraudbench":
        marker_index = next(
            (index for index, part in enumerate(parts) if part.lower() in {"positive", "negative", "deepfake"}),
            -1,
        )
        if marker_index >= 0:
            source_class = parts[marker_index]
            domain = parts[marker_index - 1] if marker_index else "unknown"
            tail = parts[marker_index + 1:]
            review_group = _review_part(tail)
            subclasses.update({
                "dataset_profile": "FraudBench",
                "domain": domain,
                "source_class": source_class,
                "generation_type": "ai_edit" if source_class.lower() == "deepfake" else "original",
            })
            if review_group:
                subclasses["review_group"] = review_group
            if source_class.lower() == "deepfake":
                generator = tail[0] if tail and not tail[0].lower().startswith("review_") else "unknown"
                subclasses["generator"] = generator
                ground_truth = "fake"
            else:
                subclasses["review_sentiment"] = source_class
                ground_truth = "real"
            stem = PurePosixPath(safe_path).stem
            source_branch = "Positive" if source_class.lower() == "deepfake" else source_class
            group_id = "/".join(filter(None, ("FraudBench", domain, source_branch, review_group, stem)))
            label_source = f"profile:fraudbench:{source_class}"

    elif profile == "rrdataset":
        transforms = {"original", "redigital", "transfer"}
        transform_index = next(
            (index for index, part in enumerate(parts[:-1]) if part.lower() in transforms),
            -1,
        )
        if transform_index >= 0:
            transform = parts[transform_index]
            source_class = parts[transform_index + 1]
            normalized_class = source_class.lower()
            subclasses.update({
                "dataset_profile": "RRDataset",
                "transform": transform,
                "source_class": source_class,
                "generation_type": "original" if normalized_class == "real" else "ai_generated",
            })
            ground_truth = "real" if normalized_class == "real" else "fake"
            stem = PurePosixPath(safe_path).stem
            content_key = re.sub(r"^(?:redigital|transfer)_", "", stem, flags=re.I)
            group_id = f"RRDataset/{normalized_class}/{content_key}"
            label_source = f"profile:rrdataset:{transform}/{source_class}"

    if ground_truth == "unlabeled":
        ground_truth, label_source = _path_label(safe_path)
        if not group_id:
            group_id = "/".join(parts) or PurePosixPath(safe_path).stem
    return {
        "groundTruth": ground_truth,
        "labelSource": label_source,
        "classPath": class_path,
        "subclasses": subclasses,
        "groupId": group_id[:1000],
    }


def _taxonomy(samples: list[dict], profile: str) -> dict:
    dimensions: dict[str, Counter] = {}
    for sample in samples:
        for key, value in (sample.get("subclasses") or {}).items():
            if key.startswith("level_") or not value:
                continue
            dimensions.setdefault(key, Counter())[str(value)] += 1
    return {
        "profile": profile,
        "dimensions": {
            key: dict(sorted(counts.items(), key=lambda item: (-item[1], item[0])))
            for key, counts in dimensions.items()
        },
        "maxDepth": max(
            (len(_path_parts(str(sample.get("relativePath") or ""))) for sample in samples),
            default=0,
        ),
    }


def _normalize_label(value: str | None) -> str:
    label = str(value or "unlabeled").strip().lower()
    return label if label in ALLOWED_LABELS else "unlabeled"


def _image_payload(data: bytes) -> tuple[str, int, int, str]:
    if not data or len(data) > MAX_UPLOAD_BYTES:
        raise ValueError("图片为空或超过 24 MB")
    try:
        with Image.open(io.BytesIO(data)) as image:
            image.verify()
        with Image.open(io.BytesIO(data)) as image:
            image_format = str(image.format or "").upper()
            width, height = image.size
    except (UnidentifiedImageError, OSError, SyntaxError, Image.DecompressionBombError) as exc:
        raise ValueError("不是可读取的图片") from exc
    if image_format not in ALLOWED_IMAGE_FORMATS:
        raise ValueError(f"不支持的图片格式：{image_format or 'unknown'}")
    if width < 32 or height < 32 or width > 8192 or height > 8192:
        raise ValueError(f"图片分辨率 {width}x{height} 超出 32-8192 像素范围")
    mime = Image.MIME.get(image_format) or mimetypes.guess_type(f"x.{image_format.lower()}")[0] or "image/octet-stream"
    suffix = {"JPEG": ".jpg", "TIFF": ".tif"}.get(image_format, f".{image_format.lower()}")
    return mime, width, height, suffix


def _pdf_images(data: bytes, source_name: str) -> list[tuple[str, bytes, str, str]]:
    try:
        import pdfplumber
    except ImportError as exc:
        raise ValueError("服务器尚未安装 PDF 图片提取组件") from exc
    images: list[tuple[str, bytes, str, str]] = []
    try:
        with pdfplumber.open(io.BytesIO(data)) as document:
            for page_number, page in enumerate(document.pages, 1):
                for image_number, item in enumerate(page.images, 1):
                    try:
                        payload = item["stream"].get_data()
                        _image_payload(payload)
                    except (KeyError, ValueError, OSError):
                        continue
                    images.append((
                        f"{Path(source_name).stem}-p{page_number}-img{image_number}",
                        payload,
                        f"pdf:{page_number}",
                        source_name,
                    ))
                    if _limit_reached(len(images), MAX_DATASET_SAMPLES):
                        return images
    except Exception as exc:
        raise ValueError("PDF 无法解析或未包含可提取图片") from exc
    return images


def _docx_images(data: bytes, source_name: str) -> list[tuple[str, bytes, str, str]]:
    try:
        from docx import Document
    except ImportError as exc:
        raise ValueError("服务器尚未安装 Word 图片提取组件") from exc
    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:
        raise ValueError("Word 文档无法解析") from exc
    images: list[tuple[str, bytes, str, str]] = []
    for index, relation in enumerate(document.part.rels.values(), 1):
        if "image" not in str(relation.target_ref):
            continue
        try:
            payload = relation.target_part.blob
            _image_payload(payload)
        except (AttributeError, ValueError, OSError):
            continue
        images.append((f"{Path(source_name).stem}-img{index}", payload, "docx", source_name))
        if _limit_reached(len(images), MAX_DATASET_SAMPLES):
            break
    return images


def _zip_relative_paths(source) -> list[str]:
    try:
        archive = zipfile.ZipFile(source if isinstance(source, Path) else _source_stream(source))
    except (OSError, zipfile.BadZipFile) as exc:
        raise ValueError("ZIP 数据集无法解析") from exc
    paths = []
    with archive:
        for member in archive.infolist():
            if member.is_dir() or member.flag_bits & 0x1:
                continue
            relative_path = _safe_relative_path(member.filename)
            suffix = Path(relative_path).suffix.lower()
            if (
                relative_path
                and relative_path != "sample"
                and suffix in {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".gif", ".tif", ".tiff", ".heic", ".heif"}
            ):
                paths.append(relative_path)
    return paths


def _iter_zip_images(source, source_name: str):
    try:
        archive = zipfile.ZipFile(source if isinstance(source, Path) else _source_stream(source))
    except (OSError, zipfile.BadZipFile) as exc:
        raise ValueError("ZIP 数据集无法解析") from exc
    with archive:
        for member in archive.infolist():
            if member.is_dir() or member.flag_bits & 0x1 or member.file_size > MAX_UPLOAD_BYTES:
                continue
            relative_path = _safe_relative_path(member.filename)
            suffix = Path(relative_path).suffix.lower()
            if suffix not in {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".gif", ".tif", ".tiff", ".heic", ".heif"}:
                continue
            try:
                payload = archive.read(member)
                _image_payload(payload)
            except (KeyError, OSError, RuntimeError, ValueError, zipfile.BadZipFile):
                continue
            yield Path(relative_path).stem, payload, f"zip:{source_name}", relative_path


def _zip_images(source, source_name: str) -> list[tuple[str, bytes, str, str]]:
    return list(_iter_zip_images(source, source_name))


def _public_https_url(value: str) -> str:
    url = str(value or "").strip()
    parsed = urlparse(url)
    if (
        parsed.scheme != "https"
        or not parsed.hostname
        or parsed.username
        or parsed.password
        or any(ord(char) <= 32 for char in url)
    ):
        raise ValueError("网页地址必须是不含账号信息的公网 HTTPS URL")
    try:
        port = parsed.port or 443
    except ValueError as exc:
        raise ValueError("网页地址端口无效") from exc
    try:
        addresses = socket.getaddrinfo(parsed.hostname, port, type=socket.SOCK_STREAM)
    except socket.gaierror as exc:
        raise ValueError("网页域名无法解析") from exc
    if not addresses:
        raise ValueError("网页域名没有可用地址")
    for entry in addresses:
        address = ipaddress.ip_address(entry[4][0])
        if not address.is_global:
            raise ValueError("网页地址不能指向内网、本机或保留地址")
    return url


def _bounded_get(session: requests.Session, url: str, *, accept_image: bool = False) -> requests.Response:
    current = _public_https_url(url)
    for _ in range(MAX_REDIRECTS + 1):
        response = session.get(
            current,
            timeout=(4, 12),
            allow_redirects=False,
            stream=True,
            headers={"User-Agent": "HuiJian-Internal-Evaluation/1.0"},
        )
        if response.status_code in {301, 302, 303, 307, 308}:
            location = response.headers.get("Location")
            response.close()
            if not location:
                raise ValueError("网页重定向缺少目标地址")
            current = _public_https_url(urljoin(current, location))
            continue
        response.raise_for_status()
        content_type = str(response.headers.get("Content-Type") or "").lower()
        if accept_image and not content_type.startswith("image/"):
            response.close()
            raise ValueError("网页资源不是图片")
        length = int(response.headers.get("Content-Length") or 0)
        if length > MAX_UPLOAD_BYTES:
            response.close()
            raise ValueError("网页资源超过 24 MB")
        return response
    raise ValueError("网页重定向次数过多")


def _read_response(response: requests.Response) -> bytes:
    chunks: list[bytes] = []
    total = 0
    try:
        for chunk in response.iter_content(64 * 1024):
            if not chunk:
                continue
            total += len(chunk)
            if total > MAX_UPLOAD_BYTES:
                raise ValueError("网页资源超过 24 MB")
            chunks.append(chunk)
    finally:
        response.close()
    return b"".join(chunks)


def _web_images(url: str) -> list[tuple[str, bytes, str, str]]:
    try:
        from bs4 import BeautifulSoup
    except ImportError as exc:
        raise ValueError("服务器尚未安装网页图片提取组件") from exc
    session = requests.Session()
    session.trust_env = False
    page_response = _bounded_get(session, url)
    page_type = str(page_response.headers.get("Content-Type") or "").lower()
    if "text/html" not in page_type:
        page_response.close()
        raise ValueError("目标地址不是 HTML 网页")
    html = _read_response(page_response)
    soup = BeautifulSoup(html, "html.parser")
    candidates: list[str] = []
    for tag in soup.find_all("img"):
        source = tag.get("src") or tag.get("data-src") or tag.get("data-original")
        if not source:
            continue
        absolute = urljoin(url, source)
        if absolute not in candidates:
            candidates.append(absolute)
        if _limit_reached(len(candidates), MAX_WEB_IMAGES):
            break
    images: list[tuple[str, bytes, str, str]] = []
    for index, source in enumerate(candidates, 1):
        try:
            payload = _read_response(_bounded_get(session, source, accept_image=True))
            _image_payload(payload)
        except (ValueError, requests.RequestException):
            continue
        images.append((f"web-image-{index}", payload, source[:500], source[:1000]))
    return images


def _iter_source(name: str, source):
    relative_path = _safe_relative_path(name)
    suffix = Path(relative_path).suffix.lower()
    if suffix == ".pdf":
        yield from _pdf_images(_read_source(source), relative_path)
        return
    if suffix == ".docx":
        yield from _docx_images(_read_source(source), relative_path)
        return
    if suffix == ".zip":
        yield from _iter_zip_images(source, relative_path)
        return
    data = _read_source(source)
    _image_payload(data)
    yield Path(relative_path).stem or "image", data, "upload", relative_path


def _extract_source(name: str, data: bytes) -> list[tuple[str, bytes, str, str]]:
    return list(_iter_source(name, data))


def create_dataset(
    uploads: list[tuple[str, object]],
    *,
    source_url: str = "",
    name: str = "",
    default_label: str = "unlabeled",
    labels: dict[str, str] | None = None,
    actor: dict | None = None,
    include_samples: bool = True,
    progress_callback=None,
    source_consumed_callback=None,
) -> dict:
    if not uploads and not source_url:
        raise ValueError("请上传图片、文件夹、ZIP、PDF、DOCX，或填写网页地址")
    if MAX_UPLOAD_FILES > 0 and len(uploads) > MAX_UPLOAD_FILES:
        raise ValueError(f"单次上传文件数不能超过 {MAX_UPLOAD_FILES}")
    ensure_schema()
    with _connect() as connection:
        usage = connection.execute(
            "SELECT COUNT(*) AS datasets, COALESCE(SUM(total_bytes),0) AS bytes FROM datasets"
        ).fetchone()
    if _limit_reached(int(usage["datasets"] or 0), MAX_STORED_DATASETS):
        raise ValueError("内部测试数据集已达到数量上限，请先删除不再使用的数据集")
    source_types: set[str] = set()
    profile_paths: list[str] = []
    for filename, source in uploads:
        relative_name = _safe_relative_path(filename)
        suffix = Path(relative_name).suffix.lower()
        source_types.add(
            "archive" if suffix == ".zip"
            else "document" if suffix in {".pdf", ".docx"}
            else "directory" if "/" in relative_name
            else "upload"
        )
        if suffix == ".zip":
            profile_paths.extend(_zip_relative_paths(source))
        else:
            profile_paths.append(relative_name)
    if source_url:
        source_types.add("web")

    dataset_id = f"ds_{uuid.uuid4().hex[:20]}"
    dataset_dir = DATA_ROOT / "datasets" / dataset_id
    dataset_dir.mkdir(parents=True, exist_ok=False)
    os.chmod(dataset_dir, 0o700)
    actor_id, actor_name = _actor_fields(actor)
    normalized_default = _normalize_label(default_label)
    label_map = {
        _safe_relative_path(key).lower(): _normalize_label(value)
        for key, value in (labels or {}).items()
    }
    seen: set[str] = set()
    samples: list[dict] = []
    total_bytes = 0
    profile = _detect_dataset_profile(profile_paths)
    estimated_samples = len(profile_paths)

    def extracted_items():
        count = 0
        for filename, source in uploads:
            try:
                for item in _iter_source(filename, source):
                    yield item
                    count += 1
                    if _limit_reached(count, MAX_DATASET_SAMPLES):
                        return
            finally:
                if source_consumed_callback:
                    source_consumed_callback(filename, source)
        if source_url:
            for item in _web_images(source_url):
                yield item
                count += 1
                if _limit_reached(count, MAX_DATASET_SAMPLES):
                    return

    try:
        for index, (sample_name, payload, source, relative_path) in enumerate(extracted_items(), 1):
            digest = hashlib.sha256(payload).hexdigest()
            if digest in seen:
                continue
            seen.add(digest)
            mime, width, height, suffix = _image_payload(payload)
            stored_name = f"{index:04d}-{digest[:16]}{suffix}"
            path = dataset_dir / stored_name
            _ensure_storage_capacity(len(payload))
            path.write_bytes(payload)
            os.chmod(path, 0o600)
            safe_sample_name = _safe_name(sample_name, f"sample-{index}")
            safe_relative_path = _safe_relative_path(relative_path, safe_sample_name)
            explicit_candidates = (
                safe_relative_path.lower(),
                Path(safe_relative_path).name.lower(),
                Path(safe_relative_path).stem.lower(),
                safe_sample_name.lower(),
            )
            explicit_label = next(
                (label_map[key] for key in explicit_candidates if key in label_map),
                None,
            )
            classification = _classify_path(safe_relative_path, profile)
            inferred_label = classification["groundTruth"]
            inferred_source = classification["labelSource"]
            if explicit_label is not None:
                label, label_source = explicit_label, "explicit"
            elif inferred_label != "unlabeled":
                label, label_source = inferred_label, inferred_source
            else:
                label = normalized_default
                label_source = "default" if label != "unlabeled" else "unresolved"
            samples.append({
                "id": f"sm_{uuid.uuid4().hex[:20]}",
                "name": safe_sample_name,
                "source": source,
                "sha256": digest,
                "mimeType": mime,
                "width": width,
                "height": height,
                "byteSize": len(payload),
                "groundTruth": label,
                "storagePath": str(path),
                "relativePath": safe_relative_path,
                "labelSource": label_source,
                "classPath": classification["classPath"],
                "subclasses": classification["subclasses"],
                "groupId": classification["groupId"],
            })
            total_bytes += len(payload)
            if progress_callback:
                progress_callback(index, estimated_samples)
        if not samples:
            raise ValueError("没有提取到符合要求的图片，或图片内容全部重复")
        created_at = _now()
        taxonomy = _taxonomy(samples, profile)
        source_type = "mixed" if len(source_types) > 1 else next(iter(source_types), "upload")
        dataset_name = _safe_name(name, f"测试集 {created_at[:10]}")
        with _connect() as connection:
            connection.execute(
                """
                INSERT INTO datasets
                    (id,name,source_type,source_name,default_label,sample_count,labeled_count,
                     total_bytes,created_at,actor_id,actor_name,profile_name,taxonomy_json)
                VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)
                """,
                (
                    dataset_id, dataset_name, source_type, source_url[:500],
                    normalized_default, len(samples),
                    sum(1 for item in samples if item["groundTruth"] != "unlabeled"),
                    total_bytes, created_at, actor_id, actor_name,
                    profile, json.dumps(taxonomy, ensure_ascii=False),
                ),
            )
            connection.executemany(
                """
                INSERT INTO samples
                    (id,dataset_id,name,source,sha256,mime_type,width,height,byte_size,
                     ground_truth,storage_path,relative_path,label_source,class_path,
                     subclasses_json,group_id,created_at)
                VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                """,
                [
                    (
                        item["id"], dataset_id, item["name"], item["source"], item["sha256"],
                        item["mimeType"], item["width"], item["height"], item["byteSize"],
                        item["groundTruth"], item["storagePath"], item["relativePath"],
                        item["labelSource"], item["classPath"],
                        json.dumps(item["subclasses"], ensure_ascii=False), item["groupId"],
                        created_at,
                    )
                    for item in samples
                ],
            )
            connection.commit()
    except Exception:
        for path in dataset_dir.glob("*"):
            path.unlink(missing_ok=True)
        dataset_dir.rmdir()
        raise
    return get_dataset(dataset_id, include_samples=include_samples) or {}


def _import_dir(import_id: str) -> Path:
    if not re.fullmatch(r"imp_[0-9a-f]{20}", str(import_id or "")):
        raise ValueError("数据集上传会话无效")
    return IMPORT_ROOT / import_id


def _import_state_path(import_id: str) -> Path:
    return _import_dir(import_id) / "session.json"


def _import_db(import_id: str) -> sqlite3.Connection:
    connection = sqlite3.connect(_import_dir(import_id) / "manifest.sqlite3", timeout=30)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA journal_mode = WAL")
    connection.execute("PRAGMA busy_timeout = 30000")
    connection.executescript(
        """
        CREATE TABLE IF NOT EXISTS files (
            id TEXT PRIMARY KEY,
            relative_path TEXT NOT NULL COLLATE NOCASE UNIQUE,
            name TEXT NOT NULL,
            storage_path TEXT NOT NULL,
            byte_size INTEGER NOT NULL,
            status TEXT NOT NULL,
            inspection_json TEXT NOT NULL DEFAULT '{}',
            created_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS chunks (
            upload_id TEXT PRIMARY KEY,
            relative_path TEXT NOT NULL,
            part_path TEXT,
            next_chunk INTEGER NOT NULL,
            total_chunks INTEGER NOT NULL,
            expected_bytes INTEGER,
            byte_size INTEGER NOT NULL DEFAULT 0,
            status TEXT NOT NULL DEFAULT 'pending',
            file_id TEXT,
            updated_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS rejections (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            relative_path TEXT NOT NULL,
            message TEXT NOT NULL,
            created_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS detections (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            file_id TEXT NOT NULL UNIQUE,
            relative_path TEXT NOT NULL,
            status TEXT NOT NULL DEFAULT 'queued',
            predicted_label TEXT,
            score REAL,
            latency_ms INTEGER,
            http_status INTEGER,
            error TEXT,
            response_json TEXT,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL,
            FOREIGN KEY(file_id) REFERENCES files(id) ON DELETE CASCADE
        );
        CREATE INDEX IF NOT EXISTS idx_import_detections_status
            ON detections(status, id);
        """
    )
    chunk_columns = {
        str(row["name"])
        for row in connection.execute("PRAGMA table_info(chunks)").fetchall()
    }
    if "expected_bytes" not in chunk_columns:
        connection.execute("ALTER TABLE chunks ADD COLUMN expected_bytes INTEGER")
        connection.commit()
    rejection_index = connection.execute(
        "SELECT 1 FROM sqlite_master WHERE type='index' AND name='idx_import_rejections_path'"
    ).fetchone()
    if not rejection_index:
        connection.execute(
            """
            DELETE FROM rejections
            WHERE id NOT IN (
                SELECT MIN(id) FROM rejections GROUP BY relative_path COLLATE NOCASE
            )
            """
        )
        connection.execute(
            "CREATE UNIQUE INDEX idx_import_rejections_path ON rejections(relative_path COLLATE NOCASE)"
        )
        connection.commit()
    return connection


def _file_row(row: sqlite3.Row) -> dict:
    item = {
        "id": row["id"],
        "relativePath": row["relative_path"],
        "name": row["name"],
        "storagePath": row["storage_path"],
        "byteSize": int(row["byte_size"] or 0),
        "status": row["status"],
    }
    try:
        item.update(json.loads(row["inspection_json"] or "{}"))
    except json.JSONDecodeError:
        pass
    return item


def _load_import(import_id: str) -> dict | None:
    try:
        path = _import_state_path(import_id)
    except ValueError:
        return None
    if not path.is_file():
        return None
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return None
    return payload if isinstance(payload, dict) else None


def _save_import(payload: dict) -> None:
    import_id = str(payload.get("id") or "")
    directory = _import_dir(import_id)
    directory.mkdir(parents=True, exist_ok=True)
    os.chmod(directory, 0o700)
    payload["updatedAt"] = _now()
    target = _import_state_path(import_id)
    temporary = target.with_suffix(".json.tmp")
    temporary.write_text(
        json.dumps(payload, ensure_ascii=False, separators=(",", ":")),
        encoding="utf-8",
    )
    os.chmod(temporary, 0o600)
    os.replace(temporary, target)


def _public_import(payload: dict | None) -> dict | None:
    if not payload:
        return None
    public = {
        key: value for key, value in payload.items()
        if key not in {"actor", "model"}
    }
    with _import_db(str(payload.get("id") or "")) as connection:
        file_rows = connection.execute(
            "SELECT * FROM files ORDER BY created_at DESC,id DESC LIMIT 100"
        ).fetchall()
        file_count = int(connection.execute("SELECT COUNT(*) FROM files").fetchone()[0])
        chunk_rows = connection.execute(
            "SELECT upload_id,relative_path,next_chunk,total_chunks,expected_bytes,byte_size FROM chunks WHERE status = 'pending' ORDER BY updated_at"
        ).fetchall()
        rejection_rows = connection.execute(
            "SELECT relative_path,message,created_at FROM rejections ORDER BY id DESC LIMIT 200"
        ).fetchall()
        rejection_count = int(connection.execute("SELECT COUNT(*) FROM rejections").fetchone()[0])
        detection_rows = connection.execute(
            """
            SELECT relative_path,status,predicted_label,score,latency_ms,error,updated_at
            FROM detections ORDER BY id DESC LIMIT 20
            """
        ).fetchall()
        detection_counts = {
            str(row["status"]): int(row["count"] or 0)
            for row in connection.execute(
                "SELECT status,COUNT(*) AS count FROM detections GROUP BY status"
            ).fetchall()
        }
    public["files"] = [
        {key: value for key, value in _file_row(row).items() if key != "storagePath"}
        for row in reversed(file_rows)
    ]
    public["filePreviewCount"] = len(public["files"])
    public["hasMoreFiles"] = file_count > len(public["files"])
    public["pendingChunks"] = [
        {
            "uploadId": row["upload_id"],
            "relativePath": row["relative_path"],
            "nextChunk": row["next_chunk"],
            "totalChunks": row["total_chunks"],
            "expectedBytes": row["expected_bytes"],
            "byteSize": row["byte_size"],
        }
        for row in chunk_rows
    ]
    public["rejections"] = [
        {
            "relativePath": row["relative_path"],
            "message": row["message"],
            "createdAt": row["created_at"],
        }
        for row in reversed(rejection_rows)
    ]
    public["rejectedFiles"] = rejection_count
    public["detection"] = {
        "total": sum(detection_counts.values()),
        "queued": detection_counts.get("queued", 0),
        "running": detection_counts.get("running", 0),
        "completed": detection_counts.get("success", 0) + detection_counts.get("failed", 0),
        "success": detection_counts.get("success", 0),
        "failed": detection_counts.get("failed", 0),
        "recent": [
            {
                "relativePath": row["relative_path"],
                "status": row["status"],
                "predictedLabel": row["predicted_label"],
                "score": row["score"],
                "latencyMs": row["latency_ms"],
                "error": row["error"],
                "updatedAt": row["updated_at"],
            }
            for row in reversed(detection_rows)
        ],
    }
    return public


def create_import_session(
    *,
    name: str = "",
    default_label: str = "unlabeled",
    source_url: str = "",
    expected_files: int = 0,
    expected_bytes: int = 0,
    stream_evaluation: bool = False,
    model: dict | None = None,
    concurrency: int = 1,
    actor: dict | None = None,
) -> dict:
    if expected_bytes > available_storage_bytes():
        raise ValueError("服务器磁盘剩余空间不足，无法接收该数据集")
    import_id = f"imp_{uuid.uuid4().hex[:20]}"
    actor_id, actor_name = _actor_fields(actor)
    stream_evaluation = bool(stream_evaluation)
    concurrency = max(1, min(int(concurrency or 1), MAX_EVALUATION_CONCURRENCY))
    if stream_evaluation and not model:
        raise ValueError("边上传边测试必须选择目标模型")
    now = _now()
    payload = {
        "id": import_id,
        "status": "uploading",
        "name": str(name or "").strip()[:180],
        "defaultLabel": _normalize_label(default_label),
        "sourceUrl": str(source_url or "").strip()[:1000],
        "expectedFiles": max(0, int(expected_files or 0)),
        "expectedBytes": max(0, int(expected_bytes or 0)),
        "uploadedFiles": 0,
        "uploadedBytes": 0,
        "validatedFiles": 0,
        "rejectedFiles": 0,
        "processedSamples": 0,
        "totalSamples": 0,
        "datasetId": None,
        "runId": None,
        "streamEvaluation": stream_evaluation,
        "modelId": str((model or {}).get("id") or ""),
        "modelName": str((model or {}).get("name") or (model or {}).get("id") or ""),
        "concurrency": concurrency,
        "model": dict(model or {}) if stream_evaluation else {},
        "error": "",
        "actor": {"adminId": actor_id, "username": actor_name},
        "createdAt": now,
        "updatedAt": now,
    }
    with _IMPORT_LOCK:
        payload_dir = _import_dir(import_id) / "payloads"
        payload_dir.mkdir(parents=True, exist_ok=False)
        os.chmod(payload_dir, 0o700)
        with _import_db(import_id):
            pass
        _save_import(payload)
    return _public_import(payload) or {}


def get_import_session(import_id: str) -> dict | None:
    with _IMPORT_LOCK:
        return _public_import(_load_import(import_id))


def get_import_resume_state(import_id: str, files: list[dict]) -> dict:
    """Match a reselected local folder against files already accepted by the server."""
    if len(files) > 1000:
        raise ValueError("单次最多校验 1000 个文件，请分批提交")
    requested: dict[str, tuple[str, int]] = {}
    for item in files:
        safe_path = _safe_relative_path(str(item.get("relativePath") or ""))
        byte_size = max(0, int(item.get("byteSize") or 0))
        requested[safe_path.casefold()] = (safe_path, byte_size)
    with _IMPORT_LOCK:
        payload = _load_import(import_id)
        if not payload:
            raise ValueError("数据集上传会话不存在")
        if payload.get("status") != "uploading":
            raise ValueError("当前上传会话不能继续上传")
        if not requested:
            return {"completedFiles": [], "conflicts": [], "pendingChunks": []}
        placeholders = ",".join("?" for _ in requested)
        paths = [item[0] for item in requested.values()]
        with _import_db(import_id) as connection:
            file_rows = connection.execute(
                f"SELECT relative_path,byte_size FROM files WHERE relative_path IN ({placeholders})",
                paths,
            ).fetchall()
            chunk_rows = connection.execute(
                f"""
                SELECT upload_id,relative_path,next_chunk,total_chunks,expected_bytes,byte_size
                FROM chunks
                WHERE status='pending' AND relative_path IN ({placeholders})
                """,
                paths,
            ).fetchall()
            rejection_rows = connection.execute(
                f"SELECT relative_path,message FROM rejections WHERE relative_path IN ({placeholders})",
                paths,
            ).fetchall()
    completed = []
    conflicts = []
    for row in file_rows:
        expected = requested.get(str(row["relative_path"]).casefold())
        if not expected:
            continue
        server_bytes = int(row["byte_size"] or 0)
        item = {"relativePath": expected[0], "byteSize": server_bytes}
        if server_bytes == expected[1]:
            completed.append(item)
        else:
            conflicts.append({**item, "localBytes": expected[1]})
    return {
        "completedFiles": completed,
        "conflicts": conflicts,
        "rejectedFiles": [
            {"relativePath": row["relative_path"], "message": row["message"]}
            for row in rejection_rows
        ],
        "pendingChunks": [
            {
                "uploadId": row["upload_id"],
                "relativePath": row["relative_path"],
                "nextChunk": int(row["next_chunk"] or 0),
                "totalChunks": int(row["total_chunks"] or 0),
                "expectedBytes": row["expected_bytes"],
                "byteSize": int(row["byte_size"] or 0),
            }
            for row in chunk_rows
        ],
    }


def list_import_sessions(limit: int = 20) -> list[dict]:
    if not IMPORT_ROOT.exists():
        return []
    sessions = []
    with _IMPORT_LOCK:
        for directory in IMPORT_ROOT.glob("imp_*"):
            raw = _load_import(directory.name)
            if (
                raw
                and raw.get("streamEvaluation")
                and raw.get("status") in {"uploading", "queued", "processing"}
                and directory.name not in _ACTIVE_IMPORT_DETECTIONS
            ):
                with _import_db(directory.name) as connection:
                    connection.execute(
                        "UPDATE detections SET status='queued',updated_at=? WHERE status='running'",
                        (_now(),),
                    )
                    connection.commit()
                _submit_import_detections(directory.name)
            if raw and raw.get("status") in {"queued", "processing"} and directory.name not in _ACTIVE_IMPORTS:
                raw["status"] = "queued"
                _save_import(raw)
                _submit_import(directory.name)
            payload = _public_import(raw)
            if payload:
                sessions.append(payload)
    sessions.sort(key=lambda item: str(item.get("updatedAt") or ""), reverse=True)
    return sessions[:max(1, int(limit))]


def _validate_staged_file(path: Path, relative_path: str) -> dict:
    suffix = Path(relative_path).suffix.lower()
    if suffix == ".zip":
        try:
            with zipfile.ZipFile(path) as archive:
                image_count = sum(
                    1 for item in archive.infolist()
                    if not item.is_dir()
                    and not item.flag_bits & 0x1
                    and Path(_safe_relative_path(item.filename)).suffix.lower()
                    in {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".gif", ".tif", ".tiff", ".heic", ".heif"}
                )
        except (OSError, zipfile.BadZipFile) as exc:
            raise ValueError("ZIP 文件损坏或无法读取") from exc
        if not image_count:
            raise ValueError("ZIP 中没有可检测图片")
        return {"kind": "archive", "imageCount": image_count}
    if suffix == ".pdf":
        with path.open("rb") as source:
            if source.read(5) != b"%PDF-":
                raise ValueError("PDF 文件头无效")
        return {"kind": "document"}
    if suffix == ".docx":
        if not zipfile.is_zipfile(path):
            raise ValueError("DOCX 文件损坏或无法读取")
        return {"kind": "document"}
    data = path.read_bytes()
    mime, width, height, _ = _image_payload(data)
    return {"kind": "image", "mimeType": mime, "width": width, "height": height}


def _insert_rejection(
    connection: sqlite3.Connection,
    payload: dict,
    relative_path: str,
    message: str,
) -> None:
    cursor = connection.execute(
        "INSERT OR IGNORE INTO rejections (relative_path,message,created_at) VALUES (?,?,?)",
        (relative_path[:1000], str(message)[:300], _now()),
    )
    if cursor.rowcount:
        payload["rejectedFiles"] = int(payload.get("rejectedFiles") or 0) + 1


def _queue_import_detection(
    connection: sqlite3.Connection,
    payload: dict,
    item: dict,
) -> bool:
    if not payload.get("streamEvaluation") or item.get("kind") != "image":
        return False
    now = _now()
    cursor = connection.execute(
        """
        INSERT OR IGNORE INTO detections
            (file_id,relative_path,status,created_at,updated_at)
        VALUES (?,?, 'queued', ?, ?)
        """,
        (item["id"], item["relativePath"], now, now),
    )
    return bool(cursor.rowcount)


def _claim_import_detection(import_id: str) -> dict | None:
    with _IMPORT_LOCK:
        payload = _load_import(import_id)
        if not payload or payload.get("status") not in {"uploading", "queued", "processing"}:
            return None
        with _import_db(import_id) as connection:
            row = connection.execute(
                """
                SELECT d.id,d.file_id,d.relative_path,f.name,f.storage_path,
                       f.inspection_json
                FROM detections d
                JOIN files f ON f.id = d.file_id
                WHERE d.status = 'queued' ORDER BY d.id LIMIT 1
                """
            ).fetchone()
            if not row:
                return None
            connection.execute(
                "UPDATE detections SET status='running',updated_at=? WHERE id=? AND status='queued'",
                (_now(), row["id"]),
            )
            connection.commit()
        return dict(row)


def _detect_import_file(import_id: str, item: dict) -> None:
    payload = _load_import(import_id) or {}
    model = payload.get("model") or {}
    path = Path(str(item.get("storage_path") or ""))
    try:
        inspection = json.loads(item.get("inspection_json") or "{}")
    except json.JSONDecodeError:
        inspection = {}
    try:
        if not path.is_file() or _import_dir(import_id) not in path.resolve().parents:
            raise ValueError("上传暂存文件已丢失")
        result = run_model(
            model,
            path.read_bytes(),
            str(item.get("name") or "sample"),
            str(inspection.get("mimeType") or "application/octet-stream"),
        )
    except Exception as exc:
        result = {
            "ok": False,
            "predictedLabel": "unknown",
            "score": None,
            "latencyMs": None,
            "httpStatus": None,
            "payload": {},
            "error": str(exc)[:500],
        }
    with _IMPORT_LOCK:
        if not _load_import(import_id):
            return
        with _import_db(import_id) as connection:
            connection.execute(
                """
                UPDATE detections
                SET status=?,predicted_label=?,score=?,latency_ms=?,http_status=?,
                    error=?,response_json=?,updated_at=?
                WHERE id=?
                """,
                (
                    "success" if result.get("ok") else "failed",
                    result.get("predictedLabel"), result.get("score"),
                    result.get("latencyMs"), result.get("httpStatus"),
                    str(result.get("error") or "")[:500],
                    json.dumps(result.get("payload") or {}, ensure_ascii=False)[:200000],
                    _now(), item["id"],
                ),
            )
            connection.commit()


def _execute_import_detections(import_id: str) -> None:
    try:
        payload = _load_import(import_id) or {}
        concurrency = max(
            1,
            min(int(payload.get("concurrency") or 1), MAX_EVALUATION_CONCURRENCY),
        )
        with ThreadPoolExecutor(max_workers=concurrency) as pool:
            futures = set()
            while True:
                while len(futures) < concurrency:
                    item = _claim_import_detection(import_id)
                    if not item:
                        break
                    futures.add(pool.submit(_detect_import_file, import_id, item))
                if not futures:
                    break
                done, futures = wait(futures, return_when=FIRST_COMPLETED)
                for future in done:
                    future.result()
    finally:
        resubmit = False
        with _IMPORT_LOCK:
            _ACTIVE_IMPORT_DETECTIONS.discard(import_id)
            if _load_import(import_id):
                with _import_db(import_id) as connection:
                    resubmit = bool(connection.execute(
                        "SELECT 1 FROM detections WHERE status='queued' LIMIT 1"
                    ).fetchone())
        if resubmit:
            _submit_import_detections(import_id)


def _submit_import_detections(import_id: str) -> None:
    with _IMPORT_LOCK:
        payload = _load_import(import_id)
        if (
            not payload
            or not payload.get("streamEvaluation")
            or import_id in _ACTIVE_IMPORT_DETECTIONS
        ):
            return
        _ACTIVE_IMPORT_DETECTIONS.add(import_id)
    _IMPORT_DETECTION_EXECUTOR.submit(_execute_import_detections, import_id)


def _wait_for_import_detections(import_id: str) -> None:
    while True:
        _submit_import_detections(import_id)
        with _IMPORT_LOCK:
            if not _load_import(import_id):
                raise ValueError("数据集上传会话已被删除")
            with _import_db(import_id) as connection:
                pending = int(connection.execute(
                    "SELECT COUNT(*) FROM detections WHERE status IN ('queued','running')"
                ).fetchone()[0])
        if not pending:
            return
        time.sleep(0.1)


def _stage_stream(import_id: str, relative_path: str, stream) -> dict:
    safe_path = _safe_relative_path(relative_path)
    suffix = Path(safe_path).suffix.lower()
    token = uuid.uuid4().hex
    target = _import_dir(import_id) / "payloads" / f"{token}{suffix[:12]}"
    total = 0
    try:
        with target.open("xb") as output:
            while True:
                chunk = stream.read(1024 * 1024)
                if not chunk:
                    break
                total += len(chunk)
                if suffix != ".zip" and total > MAX_UPLOAD_BYTES:
                    raise ValueError("单个图片或文档超过 24 MB")
                _ensure_storage_capacity(len(chunk))
                output.write(chunk)
        os.chmod(target, 0o600)
        if not total:
            raise ValueError("文件为空")
        inspection = _validate_staged_file(target, safe_path)
        return {
            "id": f"file_{token[:20]}",
            "name": _safe_name(safe_path),
            "relativePath": safe_path,
            "storagePath": str(target),
            "byteSize": total,
            "status": "validated",
            **inspection,
        }
    except Exception:
        target.unlink(missing_ok=True)
        raise


def add_import_files(import_id: str, uploads: list[tuple[str, object]]) -> dict:
    should_submit = False
    with _IMPORT_LOCK:
        payload = _load_import(import_id)
        if not payload:
            raise ValueError("数据集上传会话不存在")
        if payload.get("status") != "uploading":
            raise ValueError("当前上传会话不能继续接收文件")
        accepted = []
        with _import_db(import_id) as connection:
            payload["rejectedFiles"] = int(
                connection.execute("SELECT COUNT(*) FROM rejections").fetchone()[0]
            )
            for filename, source in uploads:
                relative_path = _safe_relative_path(filename)
                existing_row = connection.execute(
                    "SELECT * FROM files WHERE relative_path = ? COLLATE NOCASE",
                    (relative_path,),
                ).fetchone()
                if existing_row:
                    accepted.append({
                        **{
                            key: value for key, value in _file_row(existing_row).items()
                            if key != "storagePath"
                        },
                        "alreadyUploaded": True,
                    })
                    continue
                try:
                    item = _stage_stream(import_id, relative_path, _source_stream(source))
                except (OSError, ValueError) as exc:
                    _insert_rejection(connection, payload, relative_path, str(exc))
                    continue
                connection.execute(
                    """
                    INSERT INTO files
                        (id,relative_path,name,storage_path,byte_size,status,inspection_json,created_at)
                    VALUES (?,?,?,?,?,?,?,?)
                    """,
                    (
                        item["id"], item["relativePath"], item["name"], item["storagePath"],
                        item["byteSize"], item["status"],
                        json.dumps(
                            {
                                key: value for key, value in item.items()
                                if key not in {"id", "relativePath", "name", "storagePath", "byteSize", "status"}
                            },
                            ensure_ascii=False,
                        ),
                        _now(),
                    ),
                )
                should_submit = _queue_import_detection(connection, payload, item) or should_submit
                accepted.append({key: value for key, value in item.items() if key != "storagePath"})
                payload["uploadedFiles"] = int(payload.get("uploadedFiles") or 0) + 1
                payload["validatedFiles"] = int(payload.get("validatedFiles") or 0) + 1
                payload["uploadedBytes"] = int(payload.get("uploadedBytes") or 0) + int(item["byteSize"])
            connection.commit()
        _save_import(payload)
        result = _public_import(payload) or {}
        result["accepted"] = accepted
    if should_submit:
        _submit_import_detections(import_id)
    return result


def add_import_chunk(
    import_id: str,
    *,
    upload_id: str,
    relative_path: str,
    chunk_index: int,
    total_chunks: int,
    expected_bytes: int | None,
    chunk,
) -> dict:
    if not re.fullmatch(r"[0-9a-zA-Z_-]{8,80}", str(upload_id or "")):
        raise ValueError("分块上传标识无效")
    chunk_index = int(chunk_index)
    total_chunks = int(total_chunks)
    if chunk_index < 0 or total_chunks < 1 or chunk_index >= total_chunks:
        raise ValueError("分块序号无效")
    expected_bytes = int(expected_bytes) if expected_bytes is not None else None
    if expected_bytes is not None and expected_bytes < 0:
        raise ValueError("文件大小无效")
    data = _read_source(chunk, MAX_IMPORT_CHUNK_BYTES)
    should_submit = False
    with _IMPORT_LOCK:
        payload = _load_import(import_id)
        if not payload:
            raise ValueError("数据集上传会话不存在")
        if payload.get("status") != "uploading":
            raise ValueError("当前上传会话不能继续接收分块")
        safe_path = _safe_relative_path(relative_path)
        with _import_db(import_id) as connection:
            state = connection.execute(
                "SELECT * FROM chunks WHERE upload_id = ?", (upload_id,)
            ).fetchone()
            if state and state["status"] in {"completed", "rejected"}:
                return _public_import(payload) or {}
            existing = connection.execute(
                "SELECT id FROM files WHERE relative_path = ? COLLATE NOCASE", (safe_path,)
            ).fetchone()
            if existing:
                connection.execute(
                    """
                    INSERT OR REPLACE INTO chunks
                        (upload_id,relative_path,part_path,next_chunk,total_chunks,expected_bytes,byte_size,status,file_id,updated_at)
                    VALUES (?,?,?,?,?,?,?,?,?,?)
                    """,
                    (upload_id, safe_path, None, total_chunks, total_chunks, expected_bytes, 0, "completed", existing["id"], _now()),
                )
                connection.commit()
                return _public_import(payload) or {}
            if state is None:
                if chunk_index != 0:
                    raise ValueError("必须从第 1 个分块开始上传")
                part_path = _import_dir(import_id) / "payloads" / f"chunk-{upload_id}.part"
                connection.execute(
                    """
                    INSERT INTO chunks
                        (upload_id,relative_path,part_path,next_chunk,total_chunks,expected_bytes,byte_size,status,updated_at)
                    VALUES (?,?,?,?,?,?,?,?,?)
                    """,
                    (upload_id, safe_path, str(part_path), 0, total_chunks, expected_bytes, 0, "pending", _now()),
                )
                connection.commit()
                state = connection.execute(
                    "SELECT * FROM chunks WHERE upload_id = ?", (upload_id,)
                ).fetchone()
            if state["relative_path"] != safe_path or int(state["total_chunks"] or 0) != total_chunks:
                raise ValueError("分块文件信息与上传会话不一致")
            stored_expected = state["expected_bytes"]
            if stored_expected is not None and expected_bytes is not None and int(stored_expected) != expected_bytes:
                raise ValueError("本地文件大小与已保存断点不一致，请确认选择了原文件夹")
            if stored_expected is None and expected_bytes is not None:
                connection.execute(
                    "UPDATE chunks SET expected_bytes=?,updated_at=? WHERE upload_id=?",
                    (expected_bytes, _now(), upload_id),
                )
            expected = int(state["next_chunk"] or 0)
            if chunk_index < expected:
                return _public_import(payload) or {}
            if chunk_index != expected:
                raise ValueError(f"请先上传第 {expected + 1} 个分块")
            _ensure_storage_capacity(len(data))
            part_path = Path(state["part_path"])
            with part_path.open("ab") as output:
                output.write(data)
            os.chmod(part_path, 0o600)
            next_chunk = expected + 1
            byte_size = int(state["byte_size"] or 0) + len(data)
            payload["uploadedBytes"] = int(payload.get("uploadedBytes") or 0) + len(data)
            connection.execute(
                "UPDATE chunks SET next_chunk=?,byte_size=?,updated_at=? WHERE upload_id=?",
                (next_chunk, byte_size, _now(), upload_id),
            )
            if next_chunk == total_chunks:
                suffix = Path(safe_path).suffix.lower()
                rejection = ""
                if expected_bytes is not None and byte_size != expected_bytes:
                    rejection = "分块合并后的文件大小不一致，请重新上传该文件"
                elif suffix != ".zip" and byte_size > MAX_UPLOAD_BYTES:
                    rejection = "单个图片或文档超过 24 MB"
                final_path = part_path.with_name(f"{uuid.uuid4().hex}{suffix[:12]}")
                if not rejection:
                    os.replace(part_path, final_path)
                    try:
                        inspection = _validate_staged_file(final_path, safe_path)
                    except (OSError, ValueError) as exc:
                        rejection = str(exc)
                if rejection:
                    part_path.unlink(missing_ok=True)
                    final_path.unlink(missing_ok=True)
                    _insert_rejection(connection, payload, safe_path, rejection)
                    connection.execute(
                        "UPDATE chunks SET part_path=NULL,status='rejected',updated_at=? WHERE upload_id=?",
                        (_now(), upload_id),
                    )
                else:
                    item = {
                        "id": f"file_{uuid.uuid4().hex[:20]}",
                        "name": _safe_name(safe_path),
                        "relativePath": safe_path,
                        "storagePath": str(final_path),
                        "byteSize": byte_size,
                        "status": "validated",
                        **inspection,
                    }
                    connection.execute(
                        """
                        INSERT INTO files
                            (id,relative_path,name,storage_path,byte_size,status,inspection_json,created_at)
                        VALUES (?,?,?,?,?,?,?,?)
                        """,
                        (
                            item["id"], item["relativePath"], item["name"], item["storagePath"],
                            item["byteSize"], item["status"], json.dumps(inspection, ensure_ascii=False), _now(),
                        ),
                    )
                    should_submit = _queue_import_detection(connection, payload, item)
                    connection.execute(
                        "UPDATE chunks SET part_path=NULL,status='completed',file_id=?,updated_at=? WHERE upload_id=?",
                        (item["id"], _now(), upload_id),
                    )
                    payload["uploadedFiles"] = int(payload.get("uploadedFiles") or 0) + 1
                    payload["validatedFiles"] = int(payload.get("validatedFiles") or 0) + 1
            connection.commit()
        _save_import(payload)
        result = _public_import(payload) or {}
    if should_submit:
        _submit_import_detections(import_id)
    return result


def _create_stream_evaluation(import_id: str, dataset_id: str, payload: dict) -> dict:
    started_monotonic = time.monotonic()
    model = payload.get("model") or {}
    concurrency = max(
        1,
        min(int(payload.get("concurrency") or 1), MAX_EVALUATION_CONCURRENCY),
    )
    with _connect() as connection:
        samples = connection.execute(
            "SELECT * FROM samples WHERE dataset_id=? ORDER BY created_at,id",
            (dataset_id,),
        ).fetchall()
    with _import_db(import_id) as connection:
        cached_rows = connection.execute(
            "SELECT * FROM detections WHERE status IN ('success','failed')"
        ).fetchall()
    cached = {str(row["relative_path"]).lower(): row for row in cached_rows}
    run_id = f"eval_{uuid.uuid4().hex[:20]}"
    actor_id, actor_name = _actor_fields(payload.get("actor") or {})
    created_at = _now()
    model_snapshot = {
        "id": str(model.get("id") or ""),
        "name": str(model.get("name") or model.get("id") or ""),
        "version": str(model.get("version") or model.get("modelVersion") or ""),
        "runtime": str(model.get("runtime") or ""),
        "endpointSha256": hashlib.sha256(
            str(model.get("endpoint") or "").encode("utf-8")
        ).hexdigest(),
        "timeoutSeconds": min(max(int(model.get("timeoutSeconds") or 45), 2), 120),
    }
    with _connect() as connection:
        connection.execute(
            """
            INSERT INTO runs
                (id,kind,dataset_id,model_id,model_name,status,configuration_json,
                 completed_count,total_count,created_at,started_at,updated_at,actor_id,actor_name)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)
            """,
            (
                run_id, "evaluation", dataset_id, str(model.get("id") or ""),
                str(model.get("name") or model.get("id") or ""), "running",
                json.dumps(
                    {
                        "concurrency": concurrency,
                        "modelSnapshot": model_snapshot,
                        "streamedDuringUpload": True,
                    },
                    ensure_ascii=False,
                ),
                0, len(samples), created_at, created_at, created_at, actor_id, actor_name,
            ),
        )
        connection.commit()

    def result_from_cache(sample: sqlite3.Row, row: sqlite3.Row) -> dict:
        try:
            response_payload = json.loads(row["response_json"] or "{}")
        except json.JSONDecodeError:
            response_payload = {}
        return {
            "ok": row["status"] == "success",
            "sampleId": sample["id"],
            "groundTruth": sample["ground_truth"],
            "predictedLabel": row["predicted_label"] or "unknown",
            "score": row["score"],
            "latencyMs": row["latency_ms"],
            "httpStatus": row["http_status"],
            "payload": response_payload,
            "error": row["error"] or "",
            "subclasses": json.loads(sample["subclasses_json"] or "{}"),
            "groupId": sample["group_id"] or "",
        }

    def evaluate(sample: sqlite3.Row) -> dict:
        path = Path(sample["storage_path"])
        result = run_model(model, path.read_bytes(), sample["name"], sample["mime_type"])
        return {
            **result,
            "sampleId": sample["id"],
            "groundTruth": sample["ground_truth"],
            "subclasses": json.loads(sample["subclasses_json"] or "{}"),
            "groupId": sample["group_id"] or "",
        }

    completed: list[dict] = []
    missing: list[sqlite3.Row] = []
    for sample in samples:
        cached_row = cached.get(str(sample["relative_path"] or "").lower())
        if cached_row:
            completed.append(result_from_cache(sample, cached_row))
        else:
            missing.append(sample)
    resource_monitor = _ResourceMonitor(model)
    resource_monitor.start()
    if missing:
        with ThreadPoolExecutor(max_workers=concurrency) as pool:
            futures = {pool.submit(evaluate, sample): sample for sample in missing}
            for future in as_completed(futures):
                sample = futures[future]
                try:
                    completed.append(future.result())
                except Exception as exc:
                    completed.append({
                        "ok": False,
                        "sampleId": sample["id"],
                        "groundTruth": sample["ground_truth"],
                        "predictedLabel": "unknown",
                        "score": None,
                        "latencyMs": None,
                        "httpStatus": None,
                        "payload": {},
                        "error": str(exc)[:500],
                        "subclasses": json.loads(sample["subclasses_json"] or "{}"),
                        "groupId": sample["group_id"] or "",
                    })
                _set_run(run_id, completed_count=len(completed))
    with _connect() as connection:
        connection.executemany(
            """
            INSERT INTO results
                (run_id,sample_id,status,predicted_label,score,latency_ms,
                 http_status,error,response_json,created_at)
            VALUES (?,?,?,?,?,?,?,?,?,?)
            """,
            [
                (
                    run_id, result["sampleId"],
                    "success" if result.get("ok") else "failed",
                    result.get("predictedLabel"), result.get("score"),
                    result.get("latencyMs"), result.get("httpStatus"),
                    str(result.get("error") or "")[:500],
                    json.dumps(result.get("payload") or {}, ensure_ascii=False)[:200000],
                    _now(),
                )
                for result in completed
            ],
        )
        connection.commit()
    metrics = _evaluation_metrics(completed)
    metrics["timing"] = _timing_metrics(
        time.monotonic() - started_monotonic,
        completed,
        concurrency=concurrency,
        allow_latency_estimate=not missing,
    )
    metrics["resourceUsage"] = resource_monitor.finish()
    if not missing:
        metrics["resourceUsage"]["note"] = "样本在上传阶段已完成检测；此处资源值仅覆盖评测结果汇总。"
    _set_run(
        run_id,
        status="completed",
        completed_count=len(completed),
        metrics_json=json.dumps(metrics, ensure_ascii=False),
        finished_at=_now(),
    )
    return _run_row(run_id) or {}


def _execute_import(import_id: str) -> None:
    with _IMPORT_LOCK:
        payload = _load_import(import_id)
        if not payload or payload.get("status") != "queued":
            _ACTIVE_IMPORTS.discard(import_id)
            return
        payload["status"] = "processing"
        payload["error"] = ""
        _save_import(payload)
    streams = []
    try:
        payload = _load_import(import_id) or {}
        if payload.get("streamEvaluation"):
            _wait_for_import_detections(import_id)
        with _import_db(import_id) as connection:
            file_rows = connection.execute(
                "SELECT * FROM files ORDER BY created_at,id"
            ).fetchall()
        for row in file_rows:
            item = _file_row(row)
            path = Path(str(item.get("storagePath") or ""))
            if not path.is_file() or _import_dir(import_id) not in path.resolve().parents:
                raise ValueError(f"暂存文件丢失：{item.get('relativePath') or item.get('name')}")
            streams.append((str(item.get("relativePath") or item.get("name") or "sample"), path))

        def progress(processed: int, total: int) -> None:
            with _IMPORT_LOCK:
                current = _load_import(import_id)
                if not current:
                    return
                current["processedSamples"] = processed
                current["totalSamples"] = max(total, processed)
                _save_import(current)

        staged_paths = {filename: path for filename, path in streams}

        def consumed(filename: str, source) -> None:
            if not isinstance(source, Path):
                source.close()
            path = staged_paths.get(filename)
            if path:
                path.unlink(missing_ok=True)

        dataset = create_dataset(
            streams,
            source_url=str(payload.get("sourceUrl") or ""),
            name=str(payload.get("name") or ""),
            default_label=str(payload.get("defaultLabel") or "unlabeled"),
            actor=payload.get("actor") or {},
            include_samples=False,
            progress_callback=progress,
            source_consumed_callback=consumed,
        )
        run = (
            _create_stream_evaluation(import_id, str(dataset.get("id") or ""), payload)
            if payload.get("streamEvaluation")
            else {}
        )
        with _IMPORT_LOCK:
            current = _load_import(import_id) or payload
            current["status"] = "completed"
            current["datasetId"] = dataset.get("id")
            current["runId"] = run.get("id")
            current["processedSamples"] = int(dataset.get("sample_count") or 0)
            current["totalSamples"] = int(dataset.get("sample_count") or 0)
            _save_import(current)
        payload_dir = _import_dir(import_id) / "payloads"
        if payload_dir.exists():
            shutil.rmtree(payload_dir)
    except Exception as exc:
        with _IMPORT_LOCK:
            current = _load_import(import_id) or {"id": import_id}
            current["status"] = "failed"
            current["error"] = str(exc)[:1000]
            _save_import(current)
    finally:
        for _, source in streams:
            if not isinstance(source, Path):
                source.close()
        with _IMPORT_LOCK:
            _ACTIVE_IMPORTS.discard(import_id)


def _submit_import(import_id: str) -> None:
    with _IMPORT_LOCK:
        if import_id in _ACTIVE_IMPORTS:
            return
        _ACTIVE_IMPORTS.add(import_id)
    _EXECUTOR.submit(_execute_import, import_id)


def finalize_import(import_id: str) -> dict:
    with _IMPORT_LOCK:
        payload = _load_import(import_id)
        if not payload:
            raise ValueError("数据集上传会话不存在")
        if payload.get("status") != "uploading":
            raise ValueError("当前上传会话不能提交建库")
        with _import_db(import_id) as connection:
            pending_count = int(connection.execute(
                "SELECT COUNT(*) FROM chunks WHERE status = 'pending'"
            ).fetchone()[0])
            payload["rejectedFiles"] = int(
                connection.execute("SELECT COUNT(*) FROM rejections").fetchone()[0]
            )
            file_rows = connection.execute(
                "SELECT inspection_json FROM files"
            ).fetchall()
        if pending_count:
            raise ValueError("仍有文件分块尚未上传完成")
        if not file_rows and not payload.get("sourceUrl"):
            raise ValueError("没有通过校验的文件")
        payload["status"] = "queued"
        total_samples = 0
        for row in file_rows:
            try:
                inspection = json.loads(row["inspection_json"] or "{}")
            except json.JSONDecodeError:
                inspection = {}
            total_samples += int(inspection.get("imageCount") or 1)
        payload["totalSamples"] = total_samples
        _save_import(payload)
    _submit_import(import_id)
    return _public_import(payload) or {}


def delete_import_session(import_id: str) -> bool:
    with _IMPORT_LOCK:
        payload = _load_import(import_id)
        if not payload:
            return False
        if (
            payload.get("status") in {"queued", "processing"}
            or import_id in _ACTIVE_IMPORT_DETECTIONS
        ):
            raise ValueError("数据集正在建库，暂时不能取消")
        directory = _import_dir(import_id)
        if directory.exists():
            shutil.rmtree(directory)
        return True


def get_dataset(dataset_id: str, *, include_samples: bool = False) -> dict | None:
    with _connect() as connection:
        row = connection.execute("SELECT * FROM datasets WHERE id = ?", (dataset_id,)).fetchone()
        if not row:
            return None
        payload = _row_dict(row) or {}
        if include_samples:
            samples = connection.execute(
                """
                SELECT id,name,source,sha256,mime_type,width,height,byte_size,ground_truth,
                       relative_path,label_source,class_path,subclasses_json,group_id,created_at
                FROM samples WHERE dataset_id = ? ORDER BY created_at,id
                """,
                (dataset_id,),
            ).fetchall()
            payload["samples"] = [_row_dict(item) or {} for item in samples]
            source_counts = Counter(str(item["label_source"] or "unresolved") for item in samples)
            label_counts = Counter(str(item["ground_truth"] or "unlabeled") for item in samples)
            payload["classification"] = {
                "automaticCount": sum(
                    count for source, count in source_counts.items()
                    if source.startswith(("directory:", "profile:"))
                ),
                "explicitCount": int(source_counts.get("explicit", 0)),
                "defaultCount": int(source_counts.get("default", 0)),
                "unresolvedCount": int(label_counts.get("unlabeled", 0)),
                "labels": dict(label_counts),
            }
        return payload


def list_datasets(limit: int | None = None) -> list[dict]:
    with _connect() as connection:
        if limit is None:
            rows = connection.execute("SELECT * FROM datasets ORDER BY created_at DESC").fetchall()
        else:
            rows = connection.execute(
                "SELECT * FROM datasets ORDER BY created_at DESC LIMIT ?",
                (max(1, int(limit)),),
            ).fetchall()
    return [_row_dict(row) or {} for row in rows]


def sample_path(sample_id: str) -> tuple[Path, str, str] | None:
    with _connect() as connection:
        row = connection.execute(
            "SELECT storage_path,mime_type,name FROM samples WHERE id = ?", (sample_id,)
        ).fetchone()
    if not row:
        return None
    path = Path(row["storage_path"]).resolve()
    root = DATA_ROOT.resolve()
    if root not in path.parents or not path.is_file():
        return None
    return path, str(row["mime_type"]), str(row["name"])


def update_sample_label(sample_id: str, ground_truth: str) -> dict | None:
    label = _normalize_label(ground_truth)
    with _connect() as connection:
        row = connection.execute(
            "SELECT dataset_id FROM samples WHERE id = ?", (sample_id,)
        ).fetchone()
        if not row:
            return None
        connection.execute(
            "UPDATE samples SET ground_truth = ?, label_source = 'manual' WHERE id = ?",
            (label, sample_id),
        )
        connection.execute(
            """
            UPDATE datasets
            SET labeled_count = (
                SELECT COUNT(*) FROM samples
                WHERE dataset_id = datasets.id AND ground_truth != 'unlabeled'
            )
            WHERE id = ?
            """,
            (row["dataset_id"],),
        )
        updated = connection.execute(
            """
            SELECT id,name,source,sha256,mime_type,width,height,byte_size,ground_truth,
                   relative_path,label_source,class_path,subclasses_json,group_id,created_at
            FROM samples WHERE id = ?
            """,
            (sample_id,),
        ).fetchone()
        connection.commit()
    return _row_dict(updated)


def delete_dataset(dataset_id: str) -> bool:
    with _connect() as connection:
        row = connection.execute(
            "SELECT id FROM datasets WHERE id = ?", (dataset_id,)
        ).fetchone()
        if not row:
            return False
        active = connection.execute(
            """
            SELECT COUNT(*) AS count FROM runs
            WHERE dataset_id = ? AND status IN ('queued','running','cancel_requested')
            """,
            (dataset_id,),
        ).fetchone()
        if int(active["count"] or 0):
            raise ValueError("该数据集仍有运行中的任务")
        paths = [
            Path(item["storage_path"])
            for item in connection.execute(
                "SELECT storage_path FROM samples WHERE dataset_id = ?", (dataset_id,)
            ).fetchall()
        ]
        connection.execute("DELETE FROM datasets WHERE id = ?", (dataset_id,))
        connection.commit()
    dataset_dir = DATA_ROOT / "datasets" / dataset_id
    for path in paths:
        resolved = path.resolve()
        if DATA_ROOT.resolve() in resolved.parents:
            resolved.unlink(missing_ok=True)
    if dataset_dir.exists():
        shutil.rmtree(dataset_dir)
    return True


def _percentile(values: list[int], percentile: float) -> int | None:
    if not values:
        return None
    ordered = sorted(values)
    index = min(len(ordered) - 1, max(0, math.ceil(percentile * len(ordered)) - 1))
    return int(ordered[index])


def _finite_number(value) -> float | None:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def _numeric_summary(values: list[float]) -> dict:
    clean = [float(value) for value in values if _finite_number(value) is not None]
    if not clean:
        return {"count": 0, "min": None, "max": None, "mean": None, "stdDev": None,
                "p05": None, "p25": None, "p50": None, "p75": None, "p95": None}
    ordered = sorted(clean)

    def percentile(fraction: float) -> float:
        if len(ordered) == 1:
            return ordered[0]
        position = (len(ordered) - 1) * fraction
        lower = math.floor(position)
        upper = math.ceil(position)
        if lower == upper:
            return ordered[lower]
        return ordered[lower] + (ordered[upper] - ordered[lower]) * (position - lower)

    return {
        "count": len(clean),
        "min": round(ordered[0], 6),
        "max": round(ordered[-1], 6),
        "mean": round(statistics.fmean(clean), 6),
        "stdDev": round(statistics.pstdev(clean), 6) if len(clean) > 1 else 0.0,
        "p05": round(percentile(0.05), 6),
        "p25": round(percentile(0.25), 6),
        "p50": round(percentile(0.50), 6),
        "p75": round(percentile(0.75), 6),
        "p95": round(percentile(0.95), 6),
    }


def _histogram(values: list[float], *, bins: int = 10,
               lower: float | None = None, upper: float | None = None) -> dict:
    clean = [float(value) for value in values if _finite_number(value) is not None]
    if not clean:
        return {"count": 0, "lower": None, "upper": None, "bins": []}
    minimum = min(clean) if lower is None else float(lower)
    maximum = max(clean) if upper is None else float(upper)
    if maximum <= minimum:
        padding = max(abs(minimum) * 0.05, 0.5)
        minimum -= padding
        maximum += padding
    width = (maximum - minimum) / max(1, bins)
    counts = [0] * max(1, bins)
    for value in clean:
        index = int((value - minimum) / width)
        index = max(0, min(len(counts) - 1, index))
        counts[index] += 1
    return {
        "count": len(clean),
        "lower": round(minimum, 6),
        "upper": round(maximum, 6),
        "bins": [
            {
                "lower": round(minimum + index * width, 6),
                "upper": round(minimum + (index + 1) * width, 6),
                "count": count,
            }
            for index, count in enumerate(counts)
        ],
    }


def _extract_logits(payload) -> tuple[list[float] | None, list[str]]:
    """Return model logits only when the response explicitly names them as logits."""
    candidates = {"logits", "rawlogits", "raw_logits", "classlogits", "class_logits"}
    labels: list[str] = []

    def visit(value, depth: int = 0):
        nonlocal labels
        if depth > 7:
            return None
        if isinstance(value, dict):
            for key in ("classOrder", "class_order", "labels", "classes"):
                raw_labels = value.get(key)
                if isinstance(raw_labels, list) and all(isinstance(item, str) for item in raw_labels):
                    labels = [str(item)[:80] for item in raw_labels[:32]]
            for key, item in value.items():
                normalized = str(key).replace("-", "_").lower()
                if normalized in candidates:
                    raw = item[0] if isinstance(item, list) and len(item) == 1 and isinstance(item[0], list) else item
                    if isinstance(raw, (list, tuple)) and 1 <= len(raw) <= 32:
                        numbers = [_finite_number(entry) for entry in raw]
                        if all(number is not None for number in numbers):
                            return [float(number) for number in numbers]
                    scalar = _finite_number(raw)
                    if scalar is not None:
                        return [scalar]
            for item in value.values():
                found = visit(item, depth + 1)
                if found is not None:
                    return found
        elif isinstance(value, list):
            for item in value[:32]:
                found = visit(item, depth + 1)
                if found is not None:
                    return found
        return None

    return visit(payload), labels


def _binary_auc(items: list[dict]) -> float | None:
    scored = [
        (float(item["score"]), 1 if item.get("groundTruth") == "fake" else 0)
        for item in items
        if item.get("groundTruth") in {"real", "fake"}
        and _finite_number(item.get("score")) is not None
    ]
    positives = sum(label for _, label in scored)
    negatives = len(scored) - positives
    if not positives or not negatives:
        return None
    ordered = sorted(scored, key=lambda pair: pair[0])
    rank_sum = 0.0
    index = 0
    while index < len(ordered):
        end = index + 1
        while end < len(ordered) and ordered[end][0] == ordered[index][0]:
            end += 1
        average_rank = ((index + 1) + end) / 2.0
        rank_sum += average_rank * sum(label for _, label in ordered[index:end])
        index = end
    return round((rank_sum - positives * (positives + 1) / 2) / (positives * negatives), 6)


def _process_memory_mb() -> tuple[float | None, float | None]:
    current = peak = None
    try:
        status = Path("/proc/self/status").read_text(encoding="utf-8")
        for line in status.splitlines():
            if line.startswith("VmRSS:"):
                current = float(line.split()[1]) / 1024.0
            elif line.startswith("VmHWM:"):
                peak = float(line.split()[1]) / 1024.0
    except (OSError, ValueError, IndexError):
        usage = resource.getrusage(resource.RUSAGE_SELF)
        divisor = 1024.0 if os.uname().sysname != "Darwin" else 1024.0 * 1024.0
        peak = float(usage.ru_maxrss) / divisor
    return current, peak


def _host_memory_percent() -> float | None:
    try:
        values = {}
        for line in Path("/proc/meminfo").read_text(encoding="utf-8").splitlines():
            key, raw = line.split(":", 1)
            values[key] = float(raw.strip().split()[0])
        total = values.get("MemTotal", 0)
        available = values.get("MemAvailable", 0)
        return round((total - available) / total * 100.0, 2) if total else None
    except (OSError, ValueError, IndexError):
        return None


class _ResourceMonitor:
    """Low-overhead run telemetry; process values cover the shared web worker."""

    def __init__(self, model: dict) -> None:
        self.model = dict(model or {})
        self.started = time.monotonic()
        usage = resource.getrusage(resource.RUSAGE_SELF)
        self.cpu_started = float(usage.ru_utime + usage.ru_stime)
        self.stop_event = threading.Event()
        self.samples: list[dict] = []
        self.model_samples: list[dict] = []
        self.thread = threading.Thread(target=self._run, name="testing-resource-monitor", daemon=True)

    def start(self) -> None:
        self.thread.start()

    def _sample_model(self) -> dict | None:
        health_url = str(self.model.get("healthUrl") or "").strip()
        if not health_url:
            return None
        try:
            session = requests.Session()
            session.trust_env = False
            response = session.get(health_url, timeout=1.5, allow_redirects=False)
            payload = response.json() if response.ok else {}
            if not isinstance(payload, dict):
                payload = {}
            return {
                "provider": str(payload.get("activeProvider") or ""),
                "queueDepth": _finite_number(payload.get("queueDepth")),
                "gpu": payload.get("gpu") if isinstance(payload.get("gpu"), dict) else None,
            }
        except (requests.RequestException, ValueError):
            return None
        finally:
            try:
                session.close()
            except UnboundLocalError:
                pass

    def _capture(self) -> None:
        rss, peak = _process_memory_mb()
        try:
            load1 = float(os.getloadavg()[0])
        except (AttributeError, OSError):
            load1 = None
        self.samples.append({
            "rssMb": rss,
            "peakRssMb": peak,
            "hostMemoryPercent": _host_memory_percent(),
            "load1": load1,
        })
        model_sample = self._sample_model()
        if model_sample:
            self.model_samples.append(model_sample)

    def _run(self) -> None:
        while not self.stop_event.is_set():
            self._capture()
            self.stop_event.wait(0.25)

    def finish(self) -> dict:
        self._capture()
        self.stop_event.set()
        self.thread.join(timeout=2.0)
        elapsed = max(time.monotonic() - self.started, 0.001)
        usage = resource.getrusage(resource.RUSAGE_SELF)
        cpu_seconds = max(0.0, float(usage.ru_utime + usage.ru_stime) - self.cpu_started)

        def values(key: str) -> list[float]:
            return [float(item[key]) for item in self.samples if _finite_number(item.get(key)) is not None]

        gpu_devices: dict[str, list[dict]] = {}
        for sample in self.model_samples:
            gpu = sample.get("gpu") or {}
            devices = gpu.get("devices") if isinstance(gpu.get("devices"), list) else []
            for device in devices:
                if isinstance(device, dict):
                    gpu_devices.setdefault(str(device.get("index", "?")), []).append(device)
        gpu_summary = []
        for index, device_samples in sorted(gpu_devices.items()):
            util = [_finite_number(item.get("utilizationPercent")) for item in device_samples]
            used = [_finite_number(item.get("memoryUsedMb")) for item in device_samples]
            total = next((_finite_number(item.get("memoryTotalMb")) for item in device_samples if _finite_number(item.get("memoryTotalMb")) is not None), None)
            util = [item for item in util if item is not None]
            used = [item for item in used if item is not None]
            gpu_summary.append({
                "index": index,
                "utilizationMeanPercent": round(statistics.fmean(util), 2) if util else None,
                "utilizationPeakPercent": round(max(util), 2) if util else None,
                "memoryPeakMb": round(max(used), 1) if used else None,
                "memoryTotalMb": total,
            })
        providers = [item.get("provider") for item in self.model_samples if item.get("provider")]
        queue_depths = [item.get("queueDepth") for item in self.model_samples if item.get("queueDepth") is not None]
        return {
            "available": bool(self.samples),
            "sampleCount": len(self.samples),
            "samplingIntervalSeconds": 0.25,
            "webProcess": {
                "scope": "shared_worker_process",
                "cpuSeconds": round(cpu_seconds, 3),
                "averageCpuPercent": round(cpu_seconds / elapsed * 100.0, 2),
                "rssPeakMb": round(max(values("peakRssMb") or values("rssMb")), 1) if values("peakRssMb") or values("rssMb") else None,
            },
            "host": {
                "memoryPeakPercent": round(max(values("hostMemoryPercent")), 2) if values("hostMemoryPercent") else None,
                "load1Peak": round(max(values("load1")), 2) if values("load1") else None,
            },
            "modelService": {
                "sampleCount": len(self.model_samples),
                "provider": providers[-1] if providers else "",
                "queueDepthPeak": max(queue_depths) if queue_depths else None,
                "gpus": gpu_summary,
                "telemetryAvailable": bool(gpu_summary),
            },
        }


def _prediction(payload: dict) -> tuple[str, float | None]:
    data = payload.get("data") if isinstance(payload.get("data"), dict) else payload
    label = str(
        data.get("final_label")
        or data.get("verdict")
        or data.get("label")
        or ""
    ).strip().lower()
    score = data.get("fake_percentage")
    if score is not None:
        try:
            score = float(score) / 100.0
        except (TypeError, ValueError):
            score = None
    if score is None:
        try:
            score = float(data.get("aiProbability", data.get("confidence")))
        except (TypeError, ValueError):
            score = None
    if any(token in label for token in ("highly_suspected_fake", "suspected_fake", "ai生成", "fake")):
        return "fake", score
    if any(token in label for token in ("真实", "real")):
        return "real", score
    if score is not None:
        return ("fake" if score >= 0.5 else "real"), score
    return "unknown", None


def run_model(model: dict, image: bytes, filename: str, mime_type: str) -> dict:
    endpoint = str(model.get("endpoint") or "").strip()
    if not endpoint:
        raise ValueError("模型没有配置 Endpoint")
    started = time.monotonic()
    headers = {
        "X-RealGuard-Internal-Test": "1",
        "X-RealGuard-Test-Run": uuid.uuid4().hex,
    }
    files = {
        "image_file": (filename, io.BytesIO(image), mime_type),
    }
    data = {
        "openid": "__realguard_internal_test__",
        "phone": "__internal_test__",
        "source_task_id": "internal-test",
    }
    if "/api/detect" in endpoint:
        token = (
            os.environ.get("REALGUARD_V2_INTERNAL_TOKEN")
            or os.environ.get("JIANZHEN_ACCESS_TOKEN")
            or ""
        ).strip()
        if token:
            headers["X-Jianzhen-Token"] = token
        files = {"file": (filename, io.BytesIO(image), mime_type)}
        data = {
            "fileType": "image",
            "upload_consent": "1",
            "consent_version": legal_documents.CONSENT_VERSION,
            "terms_sha256": legal_documents.TERMS.sha256,
            "privacy_sha256": legal_documents.PRIVACY.sha256,
        }
    timeout = min(max(int(model.get("timeoutSeconds") or 45), 2), 120)
    session = requests.Session()
    session.trust_env = False
    try:
        response = session.post(
            endpoint,
            headers=headers,
            files=files,
            data=data,
            timeout=timeout,
            allow_redirects=False,
        )
        latency = int((time.monotonic() - started) * 1000)
        content_type = str(response.headers.get("Content-Type") or "").split(";", 1)[0].strip().lower()
        response_text = response.text[:1000]
        try:
            payload = response.json()
        except ValueError:
            payload = None
        if not isinstance(payload, dict):
            snippet = re.sub(r"\s+", " ", response_text).strip()
            if snippet.startswith("<"):
                snippet = ""
            detail = (
                f"模型返回非 JSON 对象（HTTP {response.status_code}，"
                f"Content-Type: {content_type or 'unknown'}）"
            )
            if response.status_code == 413:
                detail = "模型或反向代理拒绝了图片：请求体过大（HTTP 413）"
            elif response.status_code in {502, 503, 504}:
                detail = f"模型网关暂不可用（HTTP {response.status_code}）"
            if snippet:
                detail = f"{detail}：{snippet[:300]}"
            return {
                "ok": False,
                "httpStatus": response.status_code,
                "latencyMs": latency,
                "predictedLabel": "unknown",
                "score": None,
                "payload": {"responseFormat": "non_json", "contentType": content_type},
                "error": detail[:500],
            }
        predicted, score = _prediction(payload)
        application_code = payload.get("code")
        application_error = application_code not in (None, 0, 200, "0", "200")
        http_ok = 200 <= response.status_code < 300
        prediction_ok = predicted in {"real", "fake"}
        ok = bool(http_ok and not application_error and prediction_ok)
        error = ""
        if not http_ok:
            error = str(payload.get("message") or payload.get("msg") or response_text)[:500]
        elif application_error:
            error = str(payload.get("message") or payload.get("msg") or f"模型业务错误 code={application_code}")[:500]
        elif not prediction_ok:
            error = "模型返回了 JSON，但缺少可解析的 real/fake 结论"
        return {
            "ok": ok,
            "httpStatus": response.status_code,
            "latencyMs": latency,
            "predictedLabel": predicted,
            "score": score,
            "payload": payload,
            "error": error,
        }
    except requests.RequestException as exc:
        return {
            "ok": False,
            "httpStatus": None,
            "latencyMs": int((time.monotonic() - started) * 1000),
            "predictedLabel": "unknown",
            "score": None,
            "payload": {},
            "error": str(exc)[:500],
        }
    finally:
        session.close()


def _run_row(run_id: str) -> dict | None:
    with _connect() as connection:
        row = connection.execute("SELECT * FROM runs WHERE id = ?", (run_id,)).fetchone()
    return _row_dict(row)


def _set_run(run_id: str, **updates) -> None:
    allowed = {
        "status", "metrics_json", "error", "completed_count", "total_count",
        "started_at", "finished_at", "updated_at",
    }
    values = {key: value for key, value in updates.items() if key in allowed}
    values["updated_at"] = _now()
    assignments = ", ".join(f"{key} = ?" for key in values)
    with _connect() as connection:
        connection.execute(
            f"UPDATE runs SET {assignments} WHERE id = ?",
            (*values.values(), run_id),
        )
        connection.commit()


def _is_cancelled(run_id: str) -> bool:
    row = _run_row(run_id)
    return bool(row and row.get("status") == "cancel_requested")


def _base_evaluation_metrics(results: list[dict], *, include_distributions: bool = True) -> dict:
    latencies = [int(item["latencyMs"]) for item in results if item.get("ok")]
    labeled = [
        item for item in results
        if item.get("groundTruth") in {"real", "fake"}
        and item.get("predictedLabel") in {"real", "fake"}
    ]
    tp = sum(item["groundTruth"] == "fake" and item["predictedLabel"] == "fake" for item in labeled)
    tn = sum(item["groundTruth"] == "real" and item["predictedLabel"] == "real" for item in labeled)
    fp = sum(item["groundTruth"] == "real" and item["predictedLabel"] == "fake" for item in labeled)
    fn = sum(item["groundTruth"] == "fake" and item["predictedLabel"] == "real" for item in labeled)
    accuracy = (tp + tn) / len(labeled) if labeled else None
    precision = tp / (tp + fp) if tp + fp else None
    recall = tp / (tp + fn) if tp + fn else None
    specificity = tn / (tn + fp) if tn + fp else None
    f1 = (
        2 * precision * recall / (precision + recall)
        if precision is not None and recall is not None and precision + recall
        else None
    )
    metrics = {
        "sampleCount": len(results),
        "successCount": sum(bool(item.get("ok")) for item in results),
        "failureCount": sum(not bool(item.get("ok")) for item in results),
        "labeledCount": len(labeled),
        "confusionMatrix": {"tp": tp, "tn": tn, "fp": fp, "fn": fn},
        "accuracy": accuracy,
        "precision": precision,
        "recall": recall,
        "specificity": specificity,
        "falsePositiveRate": fp / (fp + tn) if fp + tn else None,
        "falseNegativeRate": fn / (fn + tp) if fn + tp else None,
        "balancedAccuracy": (
            (recall + specificity) / 2
            if recall is not None and specificity is not None else None
        ),
        "f1": f1,
        "auc": _binary_auc(results),
        "latency": {
            "meanMs": round(statistics.fmean(latencies), 1) if latencies else None,
            "p50Ms": _percentile(latencies, 0.50),
            "p95Ms": _percentile(latencies, 0.95),
            "p99Ms": _percentile(latencies, 0.99),
            "maxMs": max(latencies) if latencies else None,
        },
    }
    if not include_distributions:
        return metrics

    scored = [
        item for item in results
        if bool(item.get("ok")) and _finite_number(item.get("score")) is not None
    ]
    scores = [float(item["score"]) for item in scored]
    score_distribution = {
        "available": bool(scores),
        "sampleCount": len(scores),
        "threshold": 0.5,
        "summary": _numeric_summary(scores),
        "histogram": _histogram(scores, bins=10, lower=0.0, upper=1.0),
        "byGroundTruth": {},
    }
    for label in ("real", "fake"):
        label_scores = [
            float(item["score"]) for item in scored if item.get("groundTruth") == label
        ]
        score_distribution["byGroundTruth"][label] = {
            "summary": _numeric_summary(label_scores),
            "histogram": _histogram(label_scores, bins=10, lower=0.0, upper=1.0),
        }
    labeled_scores = [
        item for item in scored if item.get("groundTruth") in {"real", "fake"}
    ]
    metrics["brierScore"] = (
        round(statistics.fmean(
            (float(item["score"]) - (1.0 if item["groundTruth"] == "fake" else 0.0)) ** 2
            for item in labeled_scores
        ), 6)
        if labeled_scores else None
    )
    metrics["scoreDistribution"] = score_distribution

    vectors: list[list[float]] = []
    class_order: list[str] = []
    for item in results:
        vector, labels = _extract_logits(item.get("payload") or {})
        if vector:
            vectors.append(vector)
            if labels and len(labels) == len(vector):
                class_order = labels
    dimensions = []
    vector_size = max((len(vector) for vector in vectors), default=0)
    for index in range(vector_size):
        values = [vector[index] for vector in vectors if len(vector) > index]
        dimensions.append({
            "index": index,
            "label": class_order[index] if index < len(class_order) else f"class_{index}",
            "summary": _numeric_summary(values),
            "histogram": _histogram(values, bins=12),
        })
    metrics["logitsDistribution"] = {
        "available": bool(dimensions),
        "sampleCount": len(vectors),
        "vectorSize": vector_size,
        "classOrder": class_order,
        "dimensions": dimensions,
        "message": "" if dimensions else "当前模型接口未返回原始 logits；概率分数不会被伪装成 logits。",
    }
    metrics["errorBreakdown"] = dict(Counter(
        str(item.get("error") or "未知错误")[:120]
        for item in results if not item.get("ok")
    ))
    return metrics


def _evaluation_metrics(results: list[dict]) -> dict:
    metrics = _base_evaluation_metrics(results)
    metrics["diagnosticsVersion"] = 2
    dimensions: dict[str, dict[str, list[dict]]] = {}
    for item in results:
        for dimension, value in (item.get("subclasses") or {}).items():
            if dimension.startswith("level_") or not value:
                continue
            dimensions.setdefault(dimension, {}).setdefault(str(value), []).append(item)
    grouped = []
    for dimension, values in sorted(dimensions.items()):
        if len(values) > 60:
            continue
        for value, items in sorted(values.items(), key=lambda pair: (-len(pair[1]), pair[0])):
            item_metrics = _base_evaluation_metrics(items, include_distributions=False)
            grouped.append({
                "dimension": dimension,
                "value": value,
                "sampleCount": item_metrics["sampleCount"],
                "successCount": item_metrics["successCount"],
                "labeledCount": item_metrics["labeledCount"],
                "accuracy": item_metrics["accuracy"],
                "precision": item_metrics["precision"],
                "recall": item_metrics["recall"],
                "f1": item_metrics["f1"],
                "confusionMatrix": item_metrics["confusionMatrix"],
            })
    metrics["groupMetrics"] = grouped
    return metrics


def _timing_metrics(
    elapsed_seconds: float,
    results: list[dict],
    *,
    concurrency: int = 1,
    allow_latency_estimate: bool = False,
) -> dict:
    recorded_elapsed = max(float(elapsed_seconds or 0), 0.0)
    successful = sum(bool(item.get("ok")) for item in results)
    latencies = [
        float(item["latencyMs"])
        for item in results if _finite_number(item.get("latencyMs")) is not None
    ]
    latency_total = sum(latencies) / 1000.0
    elapsed = recorded_elapsed
    estimated = False
    if allow_latency_estimate and latencies and recorded_elapsed * 1000.0 < max(latencies):
        elapsed = latency_total / max(1, int(concurrency or 1))
        estimated = True
    return {
        "wallTimeSeconds": round(elapsed, 3),
        "recordedWallTimeSeconds": round(recorded_elapsed, 3),
        "estimated": estimated,
        "basis": "summed_latency_divided_by_concurrency" if estimated else "run_timestamps",
        "throughputSamplesPerSecond": round(len(results) / elapsed, 4) if elapsed else None,
        "successfulThroughputPerSecond": round(successful / elapsed, 4) if elapsed else None,
        "modelLatencyTotalSeconds": round(latency_total, 3),
        "concurrencyGain": round(latency_total / elapsed, 3) if elapsed else None,
    }


def _execute_evaluation(run_id: str, model: dict, concurrency: int) -> None:
    started_monotonic = time.monotonic()
    _set_run(run_id, status="running", started_at=_now())
    resource_monitor = _ResourceMonitor(model)
    resource_monitor.start()
    with _connect() as connection:
        run = connection.execute("SELECT dataset_id FROM runs WHERE id = ?", (run_id,)).fetchone()
        samples = connection.execute(
            "SELECT * FROM samples WHERE dataset_id = ? ORDER BY created_at,id",
            (run["dataset_id"],),
        ).fetchall()
    completed: list[dict] = []

    def evaluate(row: sqlite3.Row) -> dict:
        path = Path(row["storage_path"])
        result = run_model(model, path.read_bytes(), row["name"], row["mime_type"])
        return {
            **result,
            "sampleId": row["id"],
            "groundTruth": row["ground_truth"],
            "subclasses": json.loads(row["subclasses_json"] or "{}"),
            "groupId": row["group_id"] or "",
        }

    try:
        with ThreadPoolExecutor(max_workers=concurrency) as pool:
            sample_iterator = iter(samples)
            futures: dict = {}

            def fill_pending() -> None:
                while len(futures) < concurrency:
                    try:
                        row = next(sample_iterator)
                    except StopIteration:
                        return
                    futures[pool.submit(evaluate, row)] = row

            fill_pending()
            while futures:
                if _is_cancelled(run_id):
                    for pending in futures:
                        pending.cancel()
                    break
                done, _ = wait(tuple(futures), return_when=FIRST_COMPLETED)
                for future in done:
                    row = futures.pop(future)
                    try:
                        result = future.result()
                    except Exception as exc:
                        result = {
                            "ok": False,
                            "sampleId": row["id"],
                            "groundTruth": row["ground_truth"],
                            "predictedLabel": "unknown",
                            "score": None,
                            "latencyMs": None,
                            "httpStatus": None,
                            "payload": {},
                            "error": str(exc)[:500],
                            "subclasses": json.loads(row["subclasses_json"] or "{}"),
                            "groupId": row["group_id"] or "",
                        }
                    completed.append(result)
                    with _connect() as connection:
                        connection.execute(
                            """
                            INSERT INTO results
                                (run_id,sample_id,status,predicted_label,score,latency_ms,
                                 http_status,error,response_json,created_at)
                            VALUES (?,?,?,?,?,?,?,?,?,?)
                            """,
                            (
                                run_id, result["sampleId"],
                                "success" if result["ok"] else "failed",
                                result["predictedLabel"], result["score"], result["latencyMs"],
                                result["httpStatus"], result["error"],
                                json.dumps(result["payload"], ensure_ascii=False)[:200000],
                                _now(),
                            ),
                        )
                        connection.commit()
                    _set_run(run_id, completed_count=len(completed))
                fill_pending()
        status = "cancelled" if _is_cancelled(run_id) else "completed"
        metrics = _evaluation_metrics(completed)
        metrics["timing"] = _timing_metrics(
            time.monotonic() - started_monotonic,
            completed,
            concurrency=concurrency,
        )
        metrics["resourceUsage"] = resource_monitor.finish()
        _set_run(
            run_id,
            status=status,
            metrics_json=json.dumps(metrics, ensure_ascii=False),
            finished_at=_now(),
        )
    except Exception as exc:
        resource_monitor.finish()
        _set_run(run_id, status="failed", error=str(exc)[:1000], finished_at=_now())


def create_evaluation(
    dataset_id: str,
    model: dict,
    *,
    concurrency: int = 1,
    actor: dict | None = None,
) -> dict:
    dataset = get_dataset(dataset_id)
    if not dataset:
        raise ValueError("测试数据集不存在")
    concurrency = max(1, min(int(concurrency), MAX_EVALUATION_CONCURRENCY))
    run_id = f"eval_{uuid.uuid4().hex[:20]}"
    actor_id, actor_name = _actor_fields(actor)
    created_at = _now()
    model_snapshot = {
        "id": str(model.get("id") or ""),
        "name": str(model.get("name") or model.get("id") or ""),
        "version": str(model.get("version") or model.get("modelVersion") or ""),
        "runtime": str(model.get("runtime") or ""),
        "endpointSha256": hashlib.sha256(
            str(model.get("endpoint") or "").encode("utf-8")
        ).hexdigest(),
        "timeoutSeconds": min(max(int(model.get("timeoutSeconds") or 45), 2), 120),
    }
    with _connect() as connection:
        connection.execute(
            """
            INSERT INTO runs
                (id,kind,dataset_id,model_id,model_name,status,configuration_json,
                 completed_count,total_count,created_at,updated_at,actor_id,actor_name)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)
            """,
            (
                run_id, "evaluation", dataset_id, str(model.get("id") or ""),
                str(model.get("name") or model.get("id") or ""), "queued",
                json.dumps(
                    {"concurrency": concurrency, "modelSnapshot": model_snapshot},
                    ensure_ascii=False,
                ),
                0, int(dataset.get("sample_count") or 0), created_at, created_at,
                actor_id, actor_name,
            ),
        )
        connection.commit()
    _EXECUTOR.submit(_execute_evaluation, run_id, dict(model), concurrency)
    return _run_row(run_id) or {}


def _execute_load_test(
    run_id: str,
    model: dict,
    image: bytes,
    filename: str,
    mime_type: str,
    concurrency: int,
    request_count: int,
    duration_seconds: int,
) -> None:
    started_monotonic = time.monotonic()
    _set_run(run_id, status="running", started_at=_now())
    resource_monitor = _ResourceMonitor(model)
    resource_monitor.start()
    results: list[dict] = []

    def invoke() -> dict:
        return run_model(model, image, filename, mime_type)

    try:
        with ThreadPoolExecutor(max_workers=concurrency) as pool:
            pending = set()
            submitted = 0
            while submitted < request_count:
                if _is_cancelled(run_id) or time.monotonic() - started_monotonic >= duration_seconds:
                    break
                while len(pending) < concurrency and submitted < request_count:
                    pending.add(pool.submit(invoke))
                    submitted += 1
                done = {future for future in pending if future.done()}
                if not done:
                    time.sleep(0.02)
                    continue
                for future in done:
                    pending.remove(future)
                    results.append(future.result())
                _set_run(run_id, completed_count=len(results), total_count=submitted)
            for future in as_completed(pending):
                results.append(future.result())
                _set_run(run_id, completed_count=len(results), total_count=submitted)
        elapsed = max(time.monotonic() - started_monotonic, 0.001)
        latencies = [int(item["latencyMs"]) for item in results if item.get("ok")]
        status = "cancelled" if _is_cancelled(run_id) else "completed"
        metrics = {
            "requestCount": len(results),
            "successCount": sum(bool(item.get("ok")) for item in results),
            "failureCount": sum(not bool(item.get("ok")) for item in results),
            "elapsedSeconds": round(elapsed, 3),
            "throughputRps": round(len(results) / elapsed, 3),
            "errorRate": round(
                sum(not bool(item.get("ok")) for item in results) / len(results), 4
            ) if results else None,
            "latency": {
                "meanMs": round(statistics.fmean(latencies), 1) if latencies else None,
                "p50Ms": _percentile(latencies, 0.50),
                "p95Ms": _percentile(latencies, 0.95),
                "p99Ms": _percentile(latencies, 0.99),
                "maxMs": max(latencies) if latencies else None,
            },
            "statusCodes": {
                str(code): sum(item.get("httpStatus") == code for item in results)
                for code in sorted({item.get("httpStatus") for item in results}, key=str)
            },
            "resourceUsage": resource_monitor.finish(),
        }
        _set_run(
            run_id,
            status=status,
            metrics_json=json.dumps(metrics, ensure_ascii=False),
            finished_at=_now(),
        )
    except Exception as exc:
        resource_monitor.finish()
        _set_run(run_id, status="failed", error=str(exc)[:1000], finished_at=_now())


def create_load_test(
    model: dict,
    image: bytes,
    filename: str,
    mime_type: str,
    *,
    concurrency: int,
    request_count: int,
    duration_seconds: int,
    actor: dict | None = None,
) -> dict:
    _image_payload(image)
    concurrency = max(1, min(int(concurrency), MAX_LOAD_CONCURRENCY))
    request_count = max(1, min(int(request_count), MAX_LOAD_REQUESTS))
    duration_seconds = max(1, min(int(duration_seconds), MAX_LOAD_DURATION_SECONDS))
    run_id = f"load_{uuid.uuid4().hex[:20]}"
    actor_id, actor_name = _actor_fields(actor)
    configuration = {
        "concurrency": concurrency,
        "requestCount": request_count,
        "durationSeconds": duration_seconds,
        "modelSnapshot": {
            "id": str(model.get("id") or ""),
            "name": str(model.get("name") or model.get("id") or ""),
            "version": str(model.get("version") or model.get("modelVersion") or ""),
            "runtime": str(model.get("runtime") or ""),
            "endpointSha256": hashlib.sha256(
                str(model.get("endpoint") or "").encode("utf-8")
            ).hexdigest(),
        },
        "hardLimits": {
            "maxConcurrency": MAX_LOAD_CONCURRENCY,
            "maxRequests": MAX_LOAD_REQUESTS,
            "maxDurationSeconds": MAX_LOAD_DURATION_SECONDS,
        },
    }
    created_at = _now()
    with _connect() as connection:
        connection.execute(
            """
            INSERT INTO runs
                (id,kind,model_id,model_name,status,configuration_json,completed_count,
                 total_count,created_at,updated_at,actor_id,actor_name)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?)
            """,
            (
                run_id, "load_test", str(model.get("id") or ""),
                str(model.get("name") or model.get("id") or ""), "queued",
                json.dumps(configuration, ensure_ascii=False), 0, request_count,
                created_at, created_at, actor_id, actor_name,
            ),
        )
        connection.commit()
    _EXECUTOR.submit(
        _execute_load_test,
        run_id,
        dict(model),
        bytes(image),
        _safe_name(filename, "load-test.png"),
        mime_type,
        concurrency,
        request_count,
        duration_seconds,
    )
    return _run_row(run_id) or {}


def cancel_run(run_id: str) -> dict | None:
    run = _run_row(run_id)
    if not run:
        return None
    if run.get("status") in {"queued", "running"}:
        _set_run(run_id, status="cancel_requested")
    return _run_row(run_id)


def list_runs(limit: int | None = 60) -> list[dict]:
    with _connect() as connection:
        if limit is None:
            rows = connection.execute("SELECT * FROM runs ORDER BY created_at DESC").fetchall()
        else:
            rows = connection.execute(
                "SELECT * FROM runs ORDER BY created_at DESC LIMIT ?", (max(1, int(limit)),)
            ).fetchall()
    return [_row_dict(row) or {} for row in rows]


def reconcile_stale_runs(max_idle_seconds: int = 600) -> int:
    """Fail orphaned background work after a worker restart or hard timeout."""
    cutoff = (datetime.now(timezone.utc) - timedelta(seconds=max_idle_seconds)).isoformat()
    finished_at = _now()
    with _connect() as connection:
        cursor = connection.execute(
            """
            UPDATE runs
            SET status = 'failed',
                error = '后台服务重启或任务长时间无响应，运行已中止',
                finished_at = ?,
                updated_at = ?
            WHERE status IN ('queued','running','cancel_requested')
              AND updated_at < ?
            """,
            (finished_at, finished_at, cutoff),
        )
        connection.commit()
        return int(cursor.rowcount or 0)


def _elapsed_from_run(run: dict) -> float:
    try:
        started = datetime.fromisoformat(str(run.get("started_at") or run.get("created_at")))
        finished = datetime.fromisoformat(str(run.get("finished_at") or run.get("updated_at")))
        return max(0.0, (finished - started).total_seconds())
    except (TypeError, ValueError):
        return 0.0


def _backfill_evaluation_diagnostics(run: dict) -> dict:
    metrics = run.get("metrics") if isinstance(run.get("metrics"), dict) else {}
    if run.get("kind") != "evaluation" or run.get("status") not in {"completed", "cancelled"}:
        return run
    if (
        int(metrics.get("diagnosticsVersion") or 0) >= 2
        and "scoreDistribution" in metrics
        and "timing" in metrics
        and "logitsDistribution" in metrics
    ):
        return run
    results = []
    with _connect() as connection:
        rows = connection.execute(
            """
            SELECT r.status,r.predicted_label,r.score,r.latency_ms,r.http_status,r.error,
                   r.response_json,s.ground_truth,s.subclasses_json,s.group_id
            FROM results r
            LEFT JOIN samples s ON s.id = r.sample_id
            WHERE r.run_id = ? ORDER BY r.id
            """,
            (run["id"],),
        )
        for row in rows:
            try:
                payload = json.loads(row["response_json"] or "{}")
            except json.JSONDecodeError:
                payload = {}
            logits, labels = _extract_logits(payload)
            diagnostic_payload = (
                {"modelOutput": {"logits": logits, "classOrder": labels}}
                if logits else {}
            )
            try:
                subclasses = json.loads(row["subclasses_json"] or "{}")
            except json.JSONDecodeError:
                subclasses = {}
            results.append({
                "ok": row["status"] == "success",
                "groundTruth": row["ground_truth"],
                "predictedLabel": row["predicted_label"],
                "score": row["score"],
                "latencyMs": row["latency_ms"],
                "httpStatus": row["http_status"],
                "error": row["error"] or "",
                "payload": diagnostic_payload,
                "subclasses": subclasses,
                "groupId": row["group_id"] or "",
            })
    recomputed = _evaluation_metrics(results)
    configuration = run.get("configuration") if isinstance(run.get("configuration"), dict) else {}
    recomputed["timing"] = _timing_metrics(
        _elapsed_from_run(run),
        results,
        concurrency=int(configuration.get("concurrency") or 1),
        allow_latency_estimate=bool(configuration.get("streamedDuringUpload")),
    )
    recomputed["resourceUsage"] = metrics.get("resourceUsage") or {
        "available": False,
        "sampleCount": 0,
        "message": "该历史任务运行时尚未启用资源采样。",
        "webProcess": {},
        "host": {},
        "modelService": {"sampleCount": 0, "gpus": [], "telemetryAvailable": False},
    }
    run["metrics"] = {**metrics, **recomputed}
    _set_run(run["id"], metrics_json=json.dumps(run["metrics"], ensure_ascii=False))
    return run


def get_run(
    run_id: str,
    *,
    include_results: bool = True,
    result_limit: int | None = 200,
) -> dict | None:
    run = _run_row(run_id)
    if not run:
        return None
    run = _backfill_evaluation_diagnostics(run)
    if include_results:
        with _connect() as connection:
            summary = connection.execute(
                """
                SELECT COUNT(*) AS count,
                       SUM(CASE WHEN status = 'success' THEN 1 ELSE 0 END) AS success_count,
                       SUM(CASE WHEN status = 'failed' THEN 1 ELSE 0 END) AS failure_count
                FROM results WHERE run_id = ?
                """,
                (run_id,),
            ).fetchone()
            limit_sql = "" if result_limit is None else " LIMIT ?"
            params = (run_id,) if result_limit is None else (run_id, max(1, int(result_limit)))
            rows = connection.execute(
                f"""
                SELECT r.id,r.sample_id,r.status,r.predicted_label,r.score,r.latency_ms,
                       r.http_status,r.error,r.response_json,r.created_at,s.name AS sample_name,
                       s.ground_truth,s.relative_path,s.label_source,s.class_path,
                       s.subclasses_json,s.group_id
                FROM results r
                LEFT JOIN samples s ON s.id = r.sample_id
                WHERE r.run_id = ? ORDER BY r.id{limit_sql}
                """,
                params,
            ).fetchall()
        run["results"] = []
        for row in rows:
            item = _row_dict(row) or {}
            response = item.pop("response", {})
            logits, labels = _extract_logits(response)
            item["modelDiagnostics"] = {
                "logits": logits,
                "classOrder": labels,
            }
            run["results"].append(item)
        run["resultSummary"] = {
            "count": int(summary["count"] or 0),
            "successCount": int(summary["success_count"] or 0),
            "failureCount": int(summary["failure_count"] or 0),
            "returnedCount": len(rows),
            "hasMore": int(summary["count"] or 0) > len(rows),
        }
    return run


def overview() -> dict:
    reconcile_stale_runs()
    datasets = list_datasets(None)
    runs = list_runs(None)
    active = [run for run in runs if run.get("status") in {"queued", "running", "cancel_requested"}]
    imports = list_import_sessions(20)
    active_imports = [
        item for item in imports if item.get("status") in {"uploading", "queued", "processing"}
    ]
    with _connect() as connection:
        totals = connection.execute(
            """
            SELECT COUNT(*) AS dataset_count, COALESCE(SUM(sample_count),0) AS sample_count,
                   COALESCE(SUM(labeled_count),0) AS labeled_count,
                   COALESCE(SUM(total_bytes),0) AS total_bytes
            FROM datasets
            """
        ).fetchone()
        run_count = connection.execute("SELECT COUNT(*) AS count FROM runs").fetchone()
    return {
        "datasets": datasets,
        "runs": runs,
        "imports": imports,
        "summary": {
            "datasetCount": int(totals["dataset_count"] or 0),
            "sampleCount": int(totals["sample_count"] or 0),
            "labeledCount": int(totals["labeled_count"] or 0),
            "storedBytes": int(totals["total_bytes"] or 0),
            "runCount": int(run_count["count"] or 0),
            "activeRunCount": len(active),
            "activeImportCount": len(active_imports),
        },
        "limits": {
            "maxSamplesPerDataset": MAX_DATASET_SAMPLES or None,
            "maxDatasetBytes": None,
            "maxExtractedDatasetBytes": None,
            "maxEvaluationConcurrency": MAX_EVALUATION_CONCURRENCY,
            "maxLoadConcurrency": MAX_LOAD_CONCURRENCY,
            "maxLoadRequests": MAX_LOAD_REQUESTS,
            "maxLoadDurationSeconds": MAX_LOAD_DURATION_SECONDS,
            "maxStoredDatasets": MAX_STORED_DATASETS or None,
            "maxStoredBytes": None,
            "availableStorageBytes": available_storage_bytes(),
            "minimumFreeStorageBytes": MIN_FREE_STORAGE_BYTES,
        },
    }
